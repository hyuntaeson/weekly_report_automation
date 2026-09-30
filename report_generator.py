#!/usr/bin/env python3
"""
Weekly report generator based on collected activity data.
"""

from __future__ import annotations

import fnmatch
import json
import os
import re
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from jinja2 import Environment, FileSystemLoader

from database import ActivityDatabase, get_week_start_end


class ReportGenerator:
    """Generate weekly Markdown and Word reports from activity data."""

    # 상세 파일/항목 목록을 보여주지 않고 건수만 집계하면 되는 액션
    COUNT_ONLY_ACTIONS = {"email_received"}

    # 사람이 읽기 어려운 노이즈성 파일(오피스 임시/잠금 파일 등)은 보고서에서 제외.
    # file_watcher.py에서도 걸러지지만, 그 전에 이미 DB에 쌓인 데이터를 위한
    # 2차 안전장치로 보고서 생성 시점에도 한 번 더 걸러낸다.
    NOISE_FILENAME_PATTERNS = [
        "~$*", "*.tmp", "*.cache", ".~lock.*", "*.crdownload", "*.part", "*.lock",
    ]
    _HEX_TEMP_NAME_RE = re.compile(r'^[0-9A-Fa-f]{8}(\.[A-Za-z0-9]+)?$')

    def __init__(self, db_path="data/activities.db", template_dir="templates"):
        self.base_dir = Path(__file__).resolve().parent
        self.db_path = (
            str((self.base_dir / db_path).resolve())
            if not Path(db_path).is_absolute()
            else db_path
        )
        self.template_dir = (
            str((self.base_dir / template_dir).resolve())
            if not Path(template_dir).is_absolute()
            else template_dir
        )
        self.environment = Environment(
            loader=FileSystemLoader(self.template_dir),
            trim_blocks=True,
            lstrip_blocks=True,
        )

    def fetch_week_activities(self, week_start, week_end):
        """주간 활동을 DB에서 읽어 노이즈 제거·중복 제거·정규화까지 적용해 반환.
        report_pipeline(LangGraph)의 collect/filter 노드와 순차 경로가 공유."""
        activities = []
        with ActivityDatabase(self.db_path) as db:
            activities = db.get_activities_by_date_range(week_start, week_end)

        activities = [a for a in activities if not self._is_noise_activity(a)]
        activities = self._dedupe_meetings(activities)
        activities = self._drop_transient_files(activities)
        activities = [self._normalize_activity_path(a) for a in activities]
        return activities

    def analyze_week_activities(self, activities):
        """LLM/임베딩이 필요한 분석 단계 (LangGraph analyze 노드와 공유)."""
        return {
            "teams_my": self._summarize_my_teams_messages(activities),
            "topics": self._cluster_topics(activities),
        }

    def compose_weekly_data(self, week_start, week_end, activities, extras=None):
        """필터링된 활동 + 분석 결과를 보고서용 weekly_data dict로 조립."""
        extras = extras or {}
        teams_my = extras.get("teams_my") or {"count": 0, "summary": ""}
        topics = extras.get("topics") or []

        # 파일 유형별/일별 통계도 노이즈 필터·중복제거가 적용된 동일 데이터 기준으로
        # 집계해야 총계와 맞는다 (과거엔 DB 원시값을 써서 숫자가 어긋났음)
        file_type_stats = [
            {"file_type": ft, "action": ac, "count": n}
            for (ft, ac), n in sorted(
                Counter(
                    (a.get("file_type") or "unknown", a.get("action") or "unknown")
                    for a in activities
                ).items(),
                key=lambda kv: (-kv[1], kv[0]),
            )
        ]
        daily_counts = self._build_daily_counts(
            week_start, week_end, activities
        )
        by_source = self._sorted_counter(
            activity.get("source", "unknown") for activity in activities
        )
        by_action = self._sorted_counter(
            activity.get("action", "unknown") for activity in activities
        )
        top_projects_or_files = self._top_paths(activities)
        ai_tool_sessions = sum(
            1 for activity in activities if activity.get("source") == "claude_code"
        )
        recent_activities = activities[-20:]
        activities_by_action = self._group_activities_by_action(activities)

        weekly_data = {
            "week_start": week_start,
            "week_end": week_end,
            "has_activities": bool(activities),
            "total_activities": len(activities),
            "by_source": by_source,
            "by_action": by_action,
            "activities_by_action": activities_by_action,
            "by_file_type": file_type_stats,
            "daily_counts": daily_counts,
            "top_projects_or_files": top_projects_or_files,
            "ai_tool_sessions": ai_tool_sessions,
            "teams_my_message_count": teams_my["count"],
            "teams_sent_count": teams_my.get("sent", teams_my["count"]),
            "teams_received_count": teams_my.get("received", 0),
            "teams_scope": teams_my.get("scope", "mine"),
            "teams_my_summary": teams_my["summary"],
            "topics": topics,
            "raw_activities": recent_activities,
            "raw_activity_total": len(activities),
            "highlights": self._build_highlights(
                len(activities),
                by_source,
                by_action,
                daily_counts,
                top_projects_or_files,
                ai_tool_sessions,
            ),
        }

        return weekly_data

    def collect_weekly_data(self, week_start=None, week_end=None):
        """주간 데이터 수집 및 집계 (순차 실행 경로).
        LangGraph 파이프라인과 동일한 단계 함수를 순서대로 호출한다."""
        resolved_week_start, resolved_week_end = self._resolve_week_range(
            week_start, week_end
        )
        activities = self.fetch_week_activities(
            resolved_week_start, resolved_week_end
        )
        extras = self.analyze_week_activities(activities)
        return self.compose_weekly_data(
            resolved_week_start, resolved_week_end, activities, extras
        )

    def generate_markdown(self, weekly_data) -> str:
        """Jinja2 템플릿으로 Markdown 문자열 생성."""
        template = self.environment.get_template("weekly_report.md.j2")
        return template.render(**weekly_data).strip() + "\n"

    def generate_word(self, weekly_data, output_path) -> str:
        """python-docx로 Word 문서 생성."""
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        try:
            document = Document(self._ensure_word_template())

            title = document.add_paragraph()
            title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            title_run = title.add_run("주간 활동 보고서")
            title_run.bold = True
            title_run.font.size = None

            period_paragraph = document.add_paragraph()
            period_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            period_paragraph.add_run(
                f"기간: {weekly_data['week_start']} ~ {weekly_data['week_end']}"
            )

            self._add_section_title(document, "요약")
            if weekly_data["has_activities"]:
                document.add_paragraph(
                    f"이번 주 총 {weekly_data['total_activities']}건의 활동이 기록되었습니다."
                )
            else:
                document.add_paragraph("이번 주 기록된 활동이 없습니다.")

            if weekly_data["ai_tool_sessions"]:
                document.add_paragraph(
                    f"AI 코딩 관련 활동은 {weekly_data['ai_tool_sessions']}건입니다."
                )

            if weekly_data.get("topics"):
                self._add_section_title(document, "주제별 작업 (AI 분석)")
                document.add_paragraph(
                    "임베딩 기반 의미 클러스터링으로 이번 주 활동을 "
                    "주제 단위로 묶었습니다."
                )
                for topic in weekly_data["topics"]:
                    line = f"• {topic['name']} ({topic['count']}건)"
                    if topic.get("examples"):
                        line += ": " + ", ".join(topic["examples"])
                    document.add_paragraph(line)

            if weekly_data.get("teams_my_message_count") or weekly_data.get(
                "teams_received_count"
            ):
                self._add_section_title(document, "Teams 내 메시지 요약")
                if weekly_data.get("teams_scope") == "all":
                    document.add_paragraph(
                        f"이번 주 Teams에서 주고받은 메시지 "
                        f"{weekly_data['teams_my_message_count']}건"
                        f"(보낸 {weekly_data['teams_sent_count']}건 / "
                        f"받은 {weekly_data['teams_received_count']}건)을 요약했습니다."
                    )
                else:
                    document.add_paragraph(
                        f"이번 주 Teams에서 내가 보낸 메시지 "
                        f"{weekly_data['teams_my_message_count']}건을 요약했습니다."
                        f"(받은 메시지 {weekly_data['teams_received_count']}건)"
                    )
                if weekly_data.get("teams_my_summary"):
                    document.add_paragraph(weekly_data["teams_my_summary"])
                else:
                    document.add_paragraph(
                        "(AI 요약을 사용할 수 없어 건수만 표시합니다.)"
                    )

            self._add_summary_table(
                document, "소스별 활동", weekly_data["by_source"], ("소스", "건수")
            )
            self._add_summary_table(
                document, "액션별 활동", weekly_data["by_action"], ("액션", "건수")
            )

            self._add_section_title(document, "액션별 상세 내역 (실제 파일 목록)")
            activities_by_action = weekly_data.get("activities_by_action") or {}
            if activities_by_action:
                for action, data in activities_by_action.items():
                    document.add_paragraph(f"{action} ({data['total']}건)").runs[
                        0
                    ].bold = True
                    if data.get("detail_hidden"):
                        document.add_paragraph("건수만 집계됩니다 (상세 목록 생략).")
                        continue
                    for item in data["files"]:
                        text = f"• [{item['timestamp']}] {item['file_path']}"
                        if item.get("summary"):
                            text += f": {item['summary']}"
                        text += f" ({item['source']})"
                        document.add_paragraph(text)
                    if data["more"] > 0:
                        document.add_paragraph(f"... 외 {data['more']}건 더")
            else:
                document.add_paragraph("상세 내역이 없습니다.")

            self._add_section_title(document, "일별 활동 추이")
            daily_table = document.add_table(rows=1, cols=2)
            daily_header = daily_table.rows[0].cells
            daily_header[0].text = "날짜"
            daily_header[1].text = "건수"
            for item in weekly_data["daily_counts"]:
                row = daily_table.add_row().cells
                row[0].text = str(item["date"])
                row[1].text = str(item["count"])

            self._add_section_title(document, "주요 활동 하이라이트")
            for highlight in weekly_data["highlights"]:
                document.add_paragraph(f"• {highlight}")

            self._add_section_title(document, "주요 작업 파일/프로젝트")
            if weekly_data["top_projects_or_files"]:
                for path, count in weekly_data["top_projects_or_files"]:
                    document.add_paragraph(f"• {path} ({count}회)")
            else:
                document.add_paragraph("집계 가능한 파일/프로젝트 정보가 없습니다.")

            self._add_section_title(document, "파일 유형별 활동")
            if weekly_data["by_file_type"]:
                file_type_table = document.add_table(rows=1, cols=3)
                header = file_type_table.rows[0].cells
                header[0].text = "파일 유형"
                header[1].text = "액션"
                header[2].text = "건수"
                for item in weekly_data["by_file_type"]:
                    row = file_type_table.add_row().cells
                    row[0].text = str(item.get("file_type", "unknown"))
                    row[1].text = str(item.get("action", "unknown"))
                    row[2].text = str(item.get("count", 0))
            else:
                document.add_paragraph("파일 유형별 집계 데이터가 없습니다.")

            self._add_section_title(document, "최근 활동")
            if weekly_data["raw_activities"]:
                for activity in reversed(weekly_data["raw_activities"]):
                    summary = (
                        f"[{activity.get('timestamp', '')}] "
                        f"{activity.get('source', 'unknown')} / {activity.get('action', 'unknown')} / "
                        f"{activity.get('file_path', '')}"
                    )
                    document.add_paragraph(f"• {summary}")
            else:
                document.add_paragraph("표시할 최근 활동이 없습니다.")

            document.save(str(output))
        except Exception as error:
            print(f"Warning: Word report fallback activated: {error}")
            output.write_text(self.generate_markdown(weekly_data), encoding="utf-8")

        return str(output)

    def save_report(
        self, weekly_data, formats=("markdown", "word"), output_dir="reports"
    ) -> dict[str, str | None]:
        """지정된 포맷으로 보고서 파일 저장."""
        normalized_formats = {fmt.lower() for fmt in formats}
        report_dir = (
            (self.base_dir / output_dir).resolve()
            if not Path(output_dir).is_absolute()
            else Path(output_dir)
        )
        report_dir.mkdir(parents=True, exist_ok=True)

        base_name = (
            f"weekly_report_{weekly_data['week_start']}_{weekly_data['week_end']}"
        )
        file_paths: dict[str, str | None] = {"markdown": None, "word": None}

        if "markdown" in normalized_formats:
            markdown_path = report_dir / f"{base_name}.md"
            markdown_content = self.generate_markdown(weekly_data)
            markdown_path.write_text(markdown_content, encoding="utf-8")
            file_paths["markdown"] = str(markdown_path)

        if "word" in normalized_formats:
            word_path = report_dir / f"{base_name}.docx"
            file_paths["word"] = self.generate_word(weekly_data, word_path)

        return file_paths

    def generate_and_save_weekly_report(
        self, week_start=None, week_end=None, formats=("markdown", "word")
    ) -> dict[str, object]:
        """전체 파이프라인 실행 및 DB 요약 저장.
        langgraph가 있으면 그래프 경로(collect→filter→analyze→aggregate→output)로
        실행하고, 없거나 실패하면 기존 순차 경로로 폴백."""
        try:
            from report_pipeline import WeeklyReportPipeline

            return WeeklyReportPipeline(self).run(
                week_start=week_start, week_end=week_end, formats=formats
            )
        except ImportError:
            pass
        except Exception as e:
            print(f"Warning: graph pipeline failed, using sequential path: {e}")

        weekly_data = self.collect_weekly_data(week_start=week_start, week_end=week_end)
        file_paths = self.save_report(weekly_data, formats=formats)

        with ActivityDatabase(self.db_path) as db:
            db.save_weekly_summary(
                weekly_data["week_start"], weekly_data["week_end"], weekly_data
            )

        return {"weekly_data": weekly_data, "file_paths": file_paths}

    def _resolve_week_range(self, week_start, week_end):
        if week_start and week_end:
            return week_start, week_end
        if week_start and not week_end:
            end = self._parse_date(week_start) + timedelta(days=6)
            return week_start, end.strftime("%Y-%m-%d")
        if week_end and not week_start:
            start = self._parse_date(week_end) - timedelta(days=6)
            return start.strftime("%Y-%m-%d"), week_end
        return get_week_start_end()

    def _drop_transient_files(self, activities):
        """같은 기간 안에 created와 deleted가 모두 기록된 파일은 저장 후
        남지 않는 일시 파일이므로 관련 활동 전부 제거한다 (DRM 임시파일 등)."""
        created_paths, deleted_paths = set(), set()
        for a in activities:
            if a.get("source") != "filesystem":
                continue
            path = (a.get("file_path") or "").strip()
            if a.get("action") == "created":
                created_paths.add(path)
            elif a.get("action") == "deleted":
                deleted_paths.add(path)
        transient = created_paths & deleted_paths
        if not transient:
            return activities
        return [
            a for a in activities
            if not (
                a.get("source") == "filesystem"
                and (a.get("file_path") or "").strip() in transient
            )
        ]

    def _normalize_activity_path(self, activity):
        """file_path가 윈도 드라이브 경로면 '/'를 '\\'로 통일.
        수집기마다 구분자가 섞여 있어 보고서 가독성이 떨어지는 것을 방지."""
        path = activity.get("file_path")
        if isinstance(path, str) and re.match(r"^[A-Za-z]:[\\/]", path):
            activity = dict(activity)
            activity["file_path"] = path.replace("/", "\\")
        return activity

    def _build_daily_counts(self, week_start, week_end, activities):
        start_date = self._parse_date(week_start)
        end_date = self._parse_date(week_end)
        counts_by_date = Counter(
            activity.get("timestamp", "")[:10]
            for activity in activities
            if activity.get("timestamp")
        )

        ordered_counts = []
        current = start_date
        while current <= end_date:
            date_key = current.strftime("%Y-%m-%d")
            ordered_counts.append(
                {"date": date_key, "count": int(counts_by_date.get(date_key, 0))}
            )
            current += timedelta(days=1)
        return ordered_counts

    def _sorted_counter(self, values: Iterable[str]) -> dict[str, int]:
        counter = Counter(value or "unknown" for value in values)
        return dict(sorted(counter.items(), key=lambda item: (-item[1], item[0])))

    def _summarize_my_teams_messages(self, activities):
        """이번 주 Teams 메시지 중 '내가 보낸' 것만 골라 LLM으로 요약.

        판별 기준: details.sender가 teams_graph_token.json의 me_display_name과
        일치하는 메시지 (신규 수집분은 details.from_me 플래그도 있지만,
        과거 데이터 호환을 위해 sender 비교를 우선 사용)."""
        # 보고서 범위는 config/teams_settings.json의 report_scope가 결정
        try:
            from teams_collector import load_teams_settings
            scope = load_teams_settings().get("report_scope", "mine")
        except Exception:
            scope = "mine"

        token_file = self.base_dir / "config" / "teams_graph_token.json"
        me_name = None
        try:
            with open(token_file, encoding="utf-8") as f:
                me_name = json.load(f).get("me_display_name")
        except (OSError, json.JSONDecodeError):
            pass
        if not me_name and scope == "mine":
            return {"count": 0, "sent": 0, "received": 0, "scope": scope, "summary": ""}

        mine_texts, all_texts = [], []
        sent_count = received_count = 0
        for activity in activities:
            if activity.get("source") != "teams":
                continue
            details = activity.get("details")
            if isinstance(details, str):
                try:
                    details = json.loads(details)
                except (json.JSONDecodeError, TypeError):
                    continue
            if not isinstance(details, dict):
                continue
            sender = details.get("sender")
            from_me = details.get("from_me") or (me_name and sender == me_name)
            text = (details.get("text") or "").strip()
            if not text:
                continue
            chat = details.get("chat") or ""
            line = f"[{chat}] {text}" if chat else text
            all_texts.append(line)
            if from_me:
                sent_count += 1
                mine_texts.append(line)
            else:
                received_count += 1

        texts = all_texts if scope == "all" else mine_texts
        if not texts:
            return {
                "count": 0, "sent": sent_count,
                "received": received_count, "scope": scope, "summary": "",
            }

        try:
            from llm_summarizer import LLMSummarizer
            summary = LLMSummarizer().summarize(
                "\n".join(texts), max_len=600, template="teams"
            )
        except Exception as e:
            print(f"Warning: Teams message summarization failed: {e}")
            summary = None
        return {
            "count": len(texts), "sent": sent_count,
            "received": received_count, "scope": scope,
            "summary": summary or "",
        }

    def _cluster_topics(self, activities):
        """임베딩 클러스터링으로 주간 활동을 주제 단위로 묶는다.
        LiteLLM 설정이 없거나 실패하면 빈 리스트 (보고서 섹션 생략)."""
        try:
            from activity_clusterer import ActivityClusterer
            clusterer = ActivityClusterer(
                str(self.base_dir / "config" / "litellm_config.json")
            )
            if not clusterer.enabled:
                return []
            return clusterer.build_topics(activities)
        except Exception as e:
            print(f"Warning: topic clustering failed: {e}")
            return []

    def _dedupe_meetings(self, activities):
        """같은 회의가 Outlook(COM)과 Teams(Graph) 양쪽으로 수집된 경우
        참석자 정보가 풍부한 Teams 쪽만 남긴다. 제목 기준으로 비교
        (타임스탬프는 Outlook=로컬, Graph=UTC라 단순비교 불가)."""
        teams_subjects = set()
        for activity in activities:
            if activity.get("source") != "teams" or activity.get("action") != "meeting":
                continue
            subject = self._meeting_subject(activity)
            if subject:
                teams_subjects.add(subject)
        if not teams_subjects:
            return activities
        return [
            a for a in activities
            if not (
                a.get("source") == "outlook"
                and a.get("action") == "meeting"
                and self._meeting_subject(a) in teams_subjects
            )
        ]

    def _meeting_subject(self, activity):
        """활동에서 회의 제목을 정규화해서 반환 (details.subject 우선,
        없으면 file_path에서 '회의:'/'Meeting:' 프리픽스 제거)."""
        details = activity.get("details")
        if isinstance(details, str):
            try:
                details = json.loads(details)
            except (TypeError, ValueError):
                details = None
        if isinstance(details, dict) and details.get("subject"):
            return str(details["subject"]).strip().lower()
        path = str(activity.get("file_path") or "")
        for prefix in ("Teams 회의:", "Teams Meeting:", "회의:", "Meeting:"):
            if path.startswith(prefix):
                return path[len(prefix):].strip().lower()
        return path.strip().lower() or None

    # 브라우저 URL 중 인증/세션 관련 URL (쿼리에 토큰·코드가 들어가 노출되면 안 됨)
    SENSITIVE_URL_MARKERS = (
        "/callback", "session_state", "access_token", "id_token", "code=",
    )

    def _is_noise_activity(self, activity):
        """오피스 임시/잠금 파일, DRM으로 깨진 파일명, 인증 콜백 URL 등
        사람이 봐도 의미 없거나 노출하면 안 되는 항목인지 판단"""
        if activity.get("source") == "browser":
            url = str(activity.get("file_path") or "").lower()
            return any(marker in url for marker in self.SENSITIVE_URL_MARKERS)
        if activity.get("source") != "filesystem":
            return False
        file_path = str(activity.get("file_path") or "")
        if ".git" in file_path.replace("\\", "/").split("/"):
            return True
        filename = os.path.basename(file_path)
        if self._HEX_TEMP_NAME_RE.match(filename):
            return True
        if "^" in filename:
            return True
        return any(
            fnmatch.fnmatch(filename, pattern)
            for pattern in self.NOISE_FILENAME_PATTERNS
        )

    def _group_activities_by_action(self, activities, limit_per_action=15):
        """액션(created/modified/deleted/moved 등)별로 실제 파일 목록을 묶는다.

        건수만 보여주는 `by_action`과 달리, 각 액션에 대해 실제로 어떤
        파일이 생성/수정/삭제/이동되었는지 파일 경로를 나열한다.
        """
        grouped: dict[str, list[dict]] = {}
        for activity in activities:
            action = activity.get("action", "unknown")
            grouped.setdefault(action, []).append(activity)

        result = {}
        for action, items in grouped.items():
            # 최신순으로 정렬
            items_sorted = sorted(
                items, key=lambda a: a.get("timestamp", ""), reverse=True
            )
            shown = [
                {
                    "file_path": item.get("file_path", ""),
                    "timestamp": item.get("timestamp", ""),
                    "source": item.get("source", "unknown"),
                    "summary": self._extract_summary(item),
                }
                for item in items_sorted[:limit_per_action]
            ]
            result[action] = {
                "files": shown,
                "total": len(items_sorted),
                "more": max(0, len(items_sorted) - limit_per_action),
                # 받은 메일처럼 건수만 필요한 액션은 상세 목록을 숨긴다
                "detail_hidden": action in self.COUNT_ONLY_ACTIONS,
            }

        # 건수 많은 액션부터
        return dict(
            sorted(result.items(), key=lambda kv: -kv[1]["total"])
        )

    def _extract_summary(self, activity):
        """활동의 details(JSON)에서 요약 텍스트를 꺼낸다."""
        details = activity.get("details")
        if not details:
            return ""
        if isinstance(details, str):
            try:
                details = json.loads(details)
            except (TypeError, ValueError):
                return ""
        if isinstance(details, dict):
            # Teams 메시지는 summary가 없고 text가 본문이므로 함께 확인
            return str(details.get("summary") or details.get("text") or "")
        return ""

    def _top_paths(self, activities, limit=5):
        counter = Counter()
        for activity in activities:
            path = (activity.get("file_path") or "").strip()
            if not path:
                continue
            counter[path] += 1
        return counter.most_common(limit)

    def _build_highlights(
        self,
        total_activities,
        by_source,
        by_action,
        daily_counts,
        top_projects_or_files,
        ai_tool_sessions,
    ):
        if total_activities == 0:
            return ["이번 주 기록된 활동이 없습니다."]

        highlights = [f"이번 주 총 활동 수는 {total_activities}건입니다."]

        if by_source:
            top_source, top_source_count = next(iter(by_source.items()))
            highlights.append(
                f"가장 많은 활동 소스는 {top_source}이며 {top_source_count}건입니다."
            )

        if by_action:
            top_action, top_action_count = next(iter(by_action.items()))
            highlights.append(
                f"가장 많이 기록된 액션은 {top_action}이며 {top_action_count}건입니다."
            )

        peak_day = max(daily_counts, key=lambda item: item["count"], default=None)
        if peak_day:
            highlights.append(
                f"활동이 가장 많았던 날은 {peak_day['date']}로 {peak_day['count']}건입니다."
            )

        if top_projects_or_files:
            top_path, top_path_count = top_projects_or_files[0]
            highlights.append(
                f"가장 자주 등장한 파일/프로젝트는 {top_path}로 {top_path_count}회 기록되었습니다."
            )

        if ai_tool_sessions:
            highlights.append(f"Claude Code 기반 AI 활동은 {ai_tool_sessions}건입니다.")

        return highlights

    def _add_summary_table(self, document, title, items, headers):
        self._add_section_title(document, title)
        if not items:
            document.add_paragraph("집계된 데이터가 없습니다.")
            return

        table = document.add_table(rows=1, cols=2)
        header = table.rows[0].cells
        header[0].text = headers[0]
        header[1].text = headers[1]

        for key, value in items.items():
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(value)

    def _parse_date(self, value):
        return datetime.strptime(value, "%Y-%m-%d")

    def _add_section_title(self, document, title):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(title)
        run.bold = True

    def _ensure_word_template(self) -> str:
        template_path = Path(self.template_dir) / "blank_report_template.docx"
        if template_path.exists():
            try:
                Document(str(template_path))
                return str(template_path)
            except Exception:
                template_path.unlink(missing_ok=True)

        template_path.parent.mkdir(parents=True, exist_ok=True)
        self._write_minimal_docx(template_path)
        return str(template_path)

    def _write_minimal_docx(self, path: Path) -> None:
        """Build a minimal valid .docx from scratch.

        `Document()` (no args) relies on python-docx's own bundled
        `default.docx` template. On this machine that bundled file has been
        corrupted into an OLE2 container (signature `D0 CF 11 E0 ...`
        instead of the `PK` zip signature a .docx needs) — almost certainly
        by a corporate DRM/document-security tool that re-wraps .docx files
        saved to disk. Since we can't rely on that file being intact,
        generate a self-contained blank docx by hand instead.
        """
        content_types = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
            "</Types>"
        )
        package_rels = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>"
        )
        document_rels = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            "</Relationships>"
        )
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p/><w:sectPr/></w:body>"
            "</w:document>"
        )
        styles_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
            '<w:name w:val="Normal"/><w:qFormat/>'
            "</w:style>"
            "</w:styles>"
        )

        with ZipFile(path, "w", ZIP_DEFLATED) as docx_zip:
            docx_zip.writestr("[Content_Types].xml", content_types)
            docx_zip.writestr("_rels/.rels", package_rels)
            docx_zip.writestr("word/document.xml", document_xml)
            docx_zip.writestr("word/_rels/document.xml.rels", document_rels)
            docx_zip.writestr("word/styles.xml", styles_xml)


def main():
    generator = ReportGenerator()
    week_end = datetime.now().strftime("%Y-%m-%d")
    week_start = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
    result = generator.generate_and_save_weekly_report(
        week_start=week_start,
        week_end=week_end,
        formats=("markdown", "word"),
    )
    raw_file_paths = result.get("file_paths", {})
    if not isinstance(raw_file_paths, dict):
        raise RuntimeError("보고서 파일 경로 정보가 올바르지 않습니다.")
    file_paths = raw_file_paths

    print("주간 보고서 생성 완료")
    for report_format, file_path in file_paths.items():
        if file_path:
            print(f"- {report_format}: {file_path}")


if __name__ == "__main__":
    main()
