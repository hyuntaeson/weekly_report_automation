#!/usr/bin/env python3
"""
Weekly report generator based on collected activity data.
"""

from __future__ import annotations

import fnmatch
import json
import os
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Pt
from jinja2 import Environment, FileSystemLoader

from weekly_report import paths
from weekly_report.storage.database import ActivityDatabase, get_week_start_end
from weekly_report.report.sections import MY_WORK_EXCLUDED, is_my_work, sections_digest, split_sequence
from weekly_report.report.sections import build_program_sections as _build_sections


class ReportGenerator:
    """Generate weekly Markdown and Word reports from activity data."""

    # 사람이 읽기 어려운 노이즈성 파일(오피스 임시/잠금 파일 등)은 보고서에서 제외.
    # file_watcher.py에서도 걸러지지만, 그 전에 이미 DB에 쌓인 데이터를 위한
    # 2차 안전장치로 보고서 생성 시점에도 한 번 더 걸러낸다.
    NOISE_FILENAME_PATTERNS = [
        "~$*", "*.tmp", "*.cache", ".~lock.*", "*.crdownload", "*.part", "*.lock",
    ]
    _HEX_TEMP_NAME_RE = re.compile(r'^[0-9A-Fa-f]{8}(\.[A-Za-z0-9]+)?$')

    def __init__(self, db_path=paths.DB_PATH, template_dir=paths.TEMPLATES_DIR):
        # 상대 경로 인자(예: output_dir="test_reports")는 프로젝트 루트 기준
        self.base_dir = Path(paths.BASE_DIR)
        self.db_path = (
            str((self.base_dir / db_path).resolve())
            if not Path(db_path).is_absolute()
            else db_path
        )
        self.template_dir = (
            str((self.base_dir / template_dir).resolve())
            if not Path(template_dir).is_absolute()
            else template_dir
        )
        self.environment = Environment(
            loader=FileSystemLoader(self.template_dir),
            trim_blocks=True,
            lstrip_blocks=True,
        )
        self.environment.filters["circled"] = self._circled
        self.environment.filters["split_seq"] = split_sequence
        self.on_progress = None  # GUI가 진행 단계를 받아 상태 표시에 씀 (선택)

    def fetch_week_activities(self, week_start, week_end):
        """주간 활동을 DB에서 읽어 노이즈 제거·중복 제거·정규화까지 적용해 반환.
        report_pipeline(LangGraph)의 collect/filter 노드와 순차 경로가 공유."""
        activities = []
        with ActivityDatabase(self.db_path) as db:
            activities = db.get_activities_by_date_range(week_start, week_end)

        activities = [a for a in activities if not self._is_noise_activity(a)]
        activities = self._dedupe_meetings(activities)
        activities = self._drop_transient_files(activities)
        activities = [self._normalize_activity_path(a) for a in activities]
        return activities

    def build_program_sections(self, activities):
        """STEP 1: 프로그램 → 항목 구조. LLM 불필요."""
        settings, me_name = self._teams_context()
        return _build_sections(activities, me_name=me_name, teams_settings=settings)

    def analyze_week_activities(self, activities, week_start=None, week_end=None):
        """LLM/임베딩이 필요한 분석 단계 (LangGraph analyze 노드와 공유).
        week_start/end가 있으면 RAG 주제 질의 섹션도 만든다 (지난 기록 검색 기준)."""
        _, me_name = self._teams_context()
        # 파트원만 수정한 공유 문서는 내 업무가 아니므로 주제 분석에서 제외
        my_activities = [a for a in activities if is_my_work(a, me_name)]
        # 서로 독립인 분석(대부분 LLM 응답 대기)은 동시에 — 차례로 돌리면 시간이 합으로 늘어남
        self._report_progress("AI 분석 중 (Teams 요약 · 주제 분류 · 핵심 요약 · 주제 질의 · 이슈·계획)")
        with ThreadPoolExecutor(max_workers=5) as pool:
            teams_my = pool.submit(self._summarize_my_teams_messages, activities)
            topics = pool.submit(self._cluster_topics, my_activities)
            week_summary = pool.submit(self._summarize_week, self.build_program_sections(activities))
            rag_sections = pool.submit(self._rag_sections, activities, week_start, week_end)
            issues_plan = pool.submit(self._issues_and_plan, activities, week_start, week_end)
            week_summary, rag_sections = week_summary.result(), rag_sections.result()
            issues_plan = issues_plan.result() or {}
            # STEP 3은 STEP 2 결과(핵심 요약·주제 질의 요약)를 '보고 내용'으로 삼는다 —
            # 이 둘이 끝나는 대로 시작해서 Teams 요약·주제 분류와 겹쳐 돌린다
            self._report_progress("AI 분석 중 (예상 질문 생성)")
            # 이슈·계획도 '보고 내용' — 상사는 이슈 상태·다음 주 일정을 자주 묻는다
            report_points = list(week_summary or []) + [
                f"{s['topic']}: {b}" for s in rag_sections for b in s["bullets"]
            ] + [f"{i['type']}: {i['content']} (상태: {i['status']})" for i in issues_plan.get("issues", [])] + [
                f"다음 주 계획: {p['text']}" for p in issues_plan.get("plans", [])
            ]
            expected_qa = self._expected_questions(activities, week_start, week_end, report_points)
            return {
                "teams_my": teams_my.result(),
                "topics": topics.result(),
                "rag_sections": rag_sections,
                "issues_plan": issues_plan,
                "week_summary": week_summary,
                "expected_qa": expected_qa,
            }

    def _report_progress(self, message):
        """GUI 상태 표시용 진행 단계 알림 (on_progress 콜백이 있을 때만)"""
        callback = getattr(self, "on_progress", None)
        if callback:
            try:
                callback(message)
            except Exception:
                pass

    def compose_weekly_data(self, week_start, week_end, activities, extras=None):
        """필터링된 활동 + 분석 결과를 보고서용 weekly_data dict로 조립."""
        extras = extras or {}
        teams_my = extras.get("teams_my") or {"count": 0, "summary": ""}
        sections = self.build_program_sections(activities)

        program_stats = [
            {
                "name": s["name"],
                "count": s["total"],
                "days": s["active_days"],
                "items": len(s["items"]),
            }
            for s in sections
        ]

        return {
            "week_start": week_start,
            "week_end": week_end,
            "has_activities": bool(sections),
            "total_activities": len(activities),
            "program_sections": sections,
            "program_stats": program_stats,
            "week_summary": extras.get("week_summary")
            or self._fallback_week_summary(sections),
            "teams_my_message_count": teams_my["count"],
            "teams_sent_count": teams_my.get("sent", teams_my["count"]),
            "teams_received_count": teams_my.get("received", 0),
            "teams_scope": teams_my.get("scope", "mine"),
            # LLM 요약이 '# 제목'을 섞어 보내면 보고서 목차가 깨지므로 굵은 글씨로 강등
            "teams_my_summary": re.sub(
                r"(?m)^#{1,6}\s*(.+?)\s*$", r"**\1**", teams_my["summary"] or ""
            ),
            "topics": extras.get("topics") or [],
            "rag_sections": extras.get("rag_sections") or [],
            "issues": (extras.get("issues_plan") or {}).get("issues") or [],
            "plans": (extras.get("issues_plan") or {}).get("plans") or [],
            "issue_evidence": (extras.get("issues_plan") or {}).get("evidence") or [],
            "expected_qa": extras.get("expected_qa") or [],
        }

    def collect_weekly_data(self, week_start=None, week_end=None):
        """주간 데이터 수집 및 집계 (순차 실행 경로).
        LangGraph 파이프라인과 동일한 단계 함수를 순서대로 호출한다."""
        resolved_week_start, resolved_week_end = self._resolve_week_range(
            week_start, week_end
        )
        activities = self.fetch_week_activities(
            resolved_week_start, resolved_week_end
        )
        extras = self.analyze_week_activities(
            activities, resolved_week_start, resolved_week_end
        )
        return self.compose_weekly_data(
            resolved_week_start, resolved_week_end, activities, extras
        )

    def generate_markdown(self, weekly_data) -> str:
        """Jinja2 템플릿으로 Markdown 문자열 생성."""
        template = self.environment.get_template("weekly_report.md.j2")
        return template.render(**weekly_data).strip() + "\n"

    def generate_word(self, weekly_data, output_path) -> str:
        """python-docx로 Word 문서 생성."""
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        try:
            document = Document(self._ensure_word_template())
            para = self._word_para

            para(document, "주간 업무 보고서", size=18, bold=True, center=True)
            para(document, f"기간: {weekly_data['week_start']} ~ {weekly_data['week_end']}", center=True)

            para(document, "STEP 1. 프로그램별 상세 활동 내역", size=15, bold=True, space_before=12)
            sections = weekly_data.get("program_sections") or []
            if not sections:
                para(document, "이번 주 기록된 활동이 없습니다.")
            for section in sections:
                para(document, f"■ {section['name']}", size=13.5, bold=True, space_before=10)
                if section.get("note"):
                    para(document, section["note"], italic=True, size=9.5)
                for number, item in enumerate(section["items"], start=1):
                    para(document, f"{self._circled(number)} {item['title']}",
                         size=12, bold=True, indent=0.4, space_before=4)
                    for line in item["lines"]:
                        seq = split_sequence(line)
                        if seq["head"]:
                            para(document, f"- {seq['head']}", indent=1.0)
                        for sub in seq["items"]:
                            para(document, sub, indent=1.6 if seq["head"] else 1.0)

            para(document, "STEP 2. 업무 요약 (AI 종합 분석)", size=15, bold=True, space_before=16)
            para(document, "이번 주 핵심 요약", size=12, bold=True, space_before=6)
            for line in weekly_data.get("week_summary") or []:
                para(document, f"• {line}", indent=0.4)

            if weekly_data.get("topics"):
                para(document, "주제별 작업", size=12, bold=True, space_before=6)
                for topic in weekly_data["topics"]:
                    line = f"• {topic['name']} ({topic['count']}건)"
                    if topic.get("examples"):
                        line += ": " + ", ".join(topic["examples"])
                    para(document, line, indent=0.4)

            for rag in weekly_data.get("rag_sections") or []:
                if rag is weekly_data["rag_sections"][0]:
                    para(document, "주제 질의 요약", size=12, bold=True, space_before=6)
                para(document, rag["topic"], bold=True, indent=0.4, space_before=4)
                for bullet in rag["bullets"]:
                    para(document, f"• {bullet}", indent=0.8)
                para(document, "근거", size=9, italic=True, indent=0.8)
                for e in rag["evidence"]:
                    para(document, f"[{e['n']}]{'(지난)' if e['past'] else ''} {e['label']}",
                         size=9, italic=True, indent=0.8)

            if weekly_data.get("teams_my_summary"):
                para(document, "Teams 메시지 요약", size=12, bold=True, space_before=6)
                para(document, weekly_data["teams_my_summary"].replace("**", ""), indent=0.4)

            stats = weekly_data.get("program_stats") or []
            if stats:
                para(document, "프로그램별 활동 통계", size=12, bold=True, space_before=6)
                table = document.add_table(rows=1, cols=4)
                for cell, text in zip(table.rows[0].cells, ("프로그램", "활동 건수", "항목 수", "활동일")):
                    cell.text = text
                for stat in stats:
                    row = table.add_row().cells
                    row[0].text = stat["name"]
                    row[1].text = str(stat["count"])
                    row[2].text = str(stat["items"])
                    row[3].text = f"{stat['days']}일"

            expected_qa = weekly_data.get("expected_qa") or []
            if expected_qa:
                para(document, "STEP 3. 예상 질문 & 답변", size=15, bold=True, space_before=16)
            for section in expected_qa:
                para(document, f"예상질문자: {section['name']}", size=12, bold=True, space_before=6)
                if section.get("style"):
                    para(document, "질문 성향: " + " · ".join(section["style"]), size=9.5, italic=True, indent=0.4)
                for number, item in enumerate(section["qa"], start=1):
                    past = f"  ↳ 지난 대화 {', '.join(item['past'])}" if item.get("past") else ""
                    para(document, f"Q{number}. {item['q']}{past}", bold=True, indent=0.4, space_before=4)
                    para(document, f"A. {item['a']}", indent=0.8)
                if section.get("past"):
                    para(document, "지난 대화: " + " · ".join(
                        f"{p['n']} {p['when']} {p['chat']} — \"{p['text']}\"" for p in section["past"]
                    ), size=9, italic=True, indent=0.4)
                if section.get("evidence"):
                    para(document, "근거", size=9, italic=True, indent=0.4)
                    for e in section["evidence"]:
                        para(document, f"[{e['n']}] {e['label']}", size=9, italic=True, indent=0.4)
            if expected_qa:
                para(document, "※ 내가 참여한 Teams 대화만 사용 · 근거로 확인되지 않는 내용은 '확인 필요'로 표시",
                     size=9, italic=True, space_before=6)

            document.save(str(output))
        except Exception as error:
            print(f"Warning: Word report fallback activated: {error}")
            output.write_text(self.generate_markdown(weekly_data), encoding="utf-8")

        return str(output)

    def save_report(
        self, weekly_data, formats=("markdown", "word"), output_dir=paths.REPORTS_DIR
    ) -> dict[str, str | None]:
        """지정된 포맷으로 보고서 파일 저장."""
        self._report_progress("보고서 파일 저장 중")
        normalized_formats = {fmt.lower() for fmt in formats}
        report_dir = (
            (self.base_dir / output_dir).resolve()
            if not Path(output_dir).is_absolute()
            else Path(output_dir)
        )
        report_dir.mkdir(parents=True, exist_ok=True)

        base_name = (
            f"weekly_report_{weekly_data['week_start']}_{weekly_data['week_end']}"
        )
        file_paths: dict[str, str | None] = {"markdown": None, "word": None}

        if "markdown" in normalized_formats:
            markdown_path = report_dir / f"{base_name}.md"
            markdown_content = self.generate_markdown(weekly_data)
            markdown_path.write_text(markdown_content, encoding="utf-8")
            file_paths["markdown"] = str(markdown_path)

        if "word" in normalized_formats:
            word_path = report_dir / f"{base_name}.docx"
            file_paths["word"] = self.generate_word(weekly_data, word_path)

        return file_paths

    def generate_and_save_weekly_report(
        self, week_start=None, week_end=None, formats=("markdown", "word")
    ) -> dict[str, object]:
        """전체 파이프라인 실행 및 DB 요약 저장.
        langgraph가 있으면 그래프 경로(collect→filter→analyze→aggregate→output)로
        실행하고, 없거나 실패하면 기존 순차 경로로 폴백."""
        try:
            from weekly_report.report.pipeline import WeeklyReportPipeline

            return WeeklyReportPipeline(self).run(
                week_start=week_start, week_end=week_end, formats=formats
            )
        except ImportError:
            pass
        except Exception as e:
            print(f"Warning: graph pipeline failed, using sequential path: {e}")

        weekly_data = self.collect_weekly_data(week_start=week_start, week_end=week_end)
        file_paths = self.save_report(weekly_data, formats=formats)

        with ActivityDatabase(self.db_path) as db:
            db.save_weekly_summary(
                weekly_data["week_start"], weekly_data["week_end"], weekly_data
            )

        return {"weekly_data": weekly_data, "file_paths": file_paths}

    def _resolve_week_range(self, week_start, week_end):
        if week_start and week_end:
            return week_start, week_end
        if week_start and not week_end:
            end = self._parse_date(week_start) + timedelta(days=6)
            return week_start, end.strftime("%Y-%m-%d")
        if week_end and not week_start:
            start = self._parse_date(week_end) - timedelta(days=6)
            return start.strftime("%Y-%m-%d"), week_end
        return get_week_start_end()

    def _drop_transient_files(self, activities):
        """같은 기간 안에 created와 deleted가 모두 기록된 파일은 저장 후
        남지 않는 일시 파일이므로 관련 활동 전부 제거한다 (DRM 임시파일 등)."""
        created_paths, deleted_paths = set(), set()
        for a in activities:
            if a.get("source") != "filesystem":
                continue
            path = (a.get("file_path") or "").strip()
            if a.get("action") == "created":
                created_paths.add(path)
            elif a.get("action") == "deleted":
                deleted_paths.add(path)
        transient = created_paths & deleted_paths
        if not transient:
            return activities
        return [
            a for a in activities
            if not (
                a.get("source") == "filesystem"
                and (a.get("file_path") or "").strip() in transient
            )
        ]

    def _normalize_activity_path(self, activity):
        """file_path가 윈도 드라이브 경로면 '/'를 '\\'로 통일.
        수집기마다 구분자가 섞여 있어 보고서 가독성이 떨어지는 것을 방지."""
        path = activity.get("file_path")
        if isinstance(path, str) and re.match(r"^[A-Za-z]:[\\/]", path):
            activity = dict(activity)
            activity["file_path"] = path.replace("/", "\\")
        return activity

    def _teams_context(self):
        """(Teams 보고 설정, 내 표시 이름). 설정 로드 실패 시 설정은 빈 dict."""
        try:
            from weekly_report.collectors.teams import load_teams_settings
            settings = load_teams_settings() or {}
        except Exception:
            settings = {}
        me_name = None
        token_file = paths.TEAMS_TOKEN
        try:
            with open(token_file, encoding="utf-8") as f:
                me_name = json.load(f).get("me_display_name")
        except (OSError, json.JSONDecodeError):
            pass
        return settings, me_name

    WEEK_SUMMARY_PROMPT = (
        "다음은 한 주간 프로그램별 업무 활동 기록이야. 리더 주간회의 보고용으로 "
        "이번 주 수행한 업무를 핵심 위주로 3~6개 불릿으로 한국어 요약해줘. "
        "각 줄은 '- '로 시작하고 한 문장으로 쓴다. 단순 웹 열람·받은 메일·잡담은 "
        "비중을 낮추고, 직접 작성·수정·결정·공유한 일을 중심으로 묶어라. 불릿만 출력."
    )

    def _summarize_week(self, sections):
        """STEP 2 핵심 요약 — STEP 1 섹션 전체를 LLM으로 요약해 불릿 목록 반환.
        LLM 비활성/실패 시 빈 리스트 (compose 단계에서 규칙 기반 요약으로 대체)."""
        digest = sections_digest(sections)  # 공유 문서 변경(파트원 작업)은 제외됨
        if not digest:
            return []
        try:
            from weekly_report.ai.llm_summarizer import LLMSummarizer, check_summary
            result = LLMSummarizer(
                paths.LITELLM_CONFIG
            ).complete(
                self.WEEK_SUMMARY_PROMPT, digest,
                max_tokens=1500, max_input=12000, temperature=0, cache=True, validate=check_summary,
            )
        except Exception as e:
            print(f"Warning: weekly summary failed: {e}")
            return []
        if not result:
            return []
        bullets = [
            line.strip().lstrip("-•*").strip()
            for line in result.splitlines()
            if line.strip().startswith(("-", "•", "*"))
        ]
        return bullets or [result.strip()]

    def _fallback_week_summary(self, sections):
        """LLM 없이 STEP 1 섹션만으로 만드는 규칙 기반 요약 (내 업무 섹션만)."""
        sections = [s for s in sections if s["name"] not in MY_WORK_EXCLUDED]
        if not sections:
            return ["이번 주 기록된 활동이 없습니다."]
        lines = []
        for section in sections:
            titles = [item["title"] for item in section["items"]]
            examples = ", ".join(titles[:3]) + (f" 외 {len(titles) - 3}건" if len(titles) > 3 else "")
            lines.append(
                f"{section['name']}: {section['active_days']}일간 {len(titles)}개 항목 "
                f"({section['total']}건) — {examples}"
            )
        return lines

    @staticmethod
    def _circled(number):
        """1 → ①, … 20 → ⑳, 이후는 (21)"""
        return chr(0x2460 + number - 1) if 1 <= number <= 20 else f"({number})"

    def _summarize_my_teams_messages(self, activities):
        """이번 주 Teams 메시지 중 '내가 보낸' 것만 골라 LLM으로 요약.

        판별 기준: details.sender가 teams_graph_token.json의 me_display_name과
        일치하는 메시지 (신규 수집분은 details.from_me 플래그도 있지만,
        과거 데이터 호환을 위해 sender 비교를 우선 사용)."""
        # 보고서 범위는 config/teams_settings.json의 report_scope가 결정
        settings, me_name = self._teams_context()
        scope = settings.get("report_scope", "mine")
        # 설정 파일이 없을 땐 LLM 필터를 쓰지 않음 (기존 동작 유지)
        work_only = bool(settings.get("work_only", True)) if settings else False
        excluded = settings.get("excluded_chats", [])
        excluded_ids = {c.get("id") for c in excluded if isinstance(c, dict)}
        excluded_titles = {c.get("title") for c in excluded if isinstance(c, dict)}

        if not me_name and scope == "mine":
            return {"count": 0, "sent": 0, "received": 0, "scope": scope, "summary": ""}

        mine_texts, all_texts = [], []
        sent_count = received_count = 0
        for activity in activities:
            if activity.get("source") != "teams":
                continue
            details = activity.get("details")
            if isinstance(details, str):
                try:
                    details = json.loads(details)
                except (json.JSONDecodeError, TypeError):
                    continue
            if not isinstance(details, dict):
                continue
            # Settings에서 제외한 채팅방은 요약 대상이 아님 (id 우선, 과거 데이터는 제목으로)
            if details.get("chat_id") in excluded_ids or details.get("chat") in excluded_titles:
                continue
            sender = details.get("sender")
            from_me = details.get("from_me") or (me_name and sender == me_name)
            text = (details.get("text") or "").strip()
            if not text:
                continue
            chat = details.get("chat") or ""
            line = f"[{chat}] {text}" if chat else text
            all_texts.append(line)
            if from_me:
                sent_count += 1
                mine_texts.append(line)
            else:
                received_count += 1

        texts = all_texts if scope == "all" else mine_texts
        if not texts:
            return {
                "count": 0, "sent": sent_count,
                "received": received_count, "scope": scope, "summary": "",
            }

        # 업무 관련 메시지만 선별 — 인사·잡담·이모티콘·비속어 제외
        if work_only:
            texts = self._filter_work_messages(texts)
            if not texts:
                return {
                    "count": 0, "sent": sent_count,
                    "received": received_count, "scope": scope, "summary": "",
                }

        try:
            from weekly_report.ai.llm_summarizer import LLMSummarizer
            summary = LLMSummarizer().summarize(
                "\n".join(texts), max_len=600, template="teams"
            )
        except Exception as e:
            print(f"Warning: Teams message summarization failed: {e}")
            summary = None
        return {
            "count": len(texts), "sent": sent_count,
            "received": received_count, "scope": scope,
            "summary": summary or "",
        }

    def _filter_work_messages(self, texts):
        """LLM으로 업무 관련 메시지만 선별.

        요청·지시·진행 상황·이슈·결정·질의응답 등 업무와 관련된 메시지의
        번호만 돌려받는다. 인사·잡담·이모티콘·비속어·업무 무관 내용은 제외.
        LLM 실패 시에는 안전하게 원본 리스트를 그대로 반환."""
        if not texts:
            return texts
        numbered = "\n".join(f"{i+1}. {t}" for i, t in enumerate(texts))
        prompt = (
            "다음은 Teams 채팅 메시지 목록이다. "
            "업무와 관련된 메시지(요청·지시·진행 상황 보고·이슈·결정·"
            "질의응답·문서/파일 관련 등)의 번호만 골라 콤마로 구분해 출력해라. "
            "인사·잡담·이모티콘·비속어·업무와 무관한 내용은 제외한다. "
            "업무 관련이 없으면 'NONE'만 출력. 번호만 출력하고 설명은 쓰지 마라."
        )
        try:
            from weekly_report.ai.llm_summarizer import LLMSummarizer
            result = LLMSummarizer().complete(prompt, numbered, max_tokens=500, temperature=0, cache=True)
            if not result or "NONE" in result.upper():
                return []
            picked = []
            for token in re.split(r"[^0-9]+", result):
                if not token:
                    continue
                idx = int(token) - 1
                if 0 <= idx < len(texts):
                    picked.append(texts[idx])
            return picked if picked else texts
        except Exception as e:
            print(f"Warning: work-message filtering failed: {e}")
            return texts

    def _cluster_topics(self, activities):
        """임베딩 클러스터링으로 주간 활동을 주제 단위로 묶는다.
        LiteLLM 설정이 없거나 실패하면 빈 리스트 (보고서 섹션 생략)."""
        try:
            from weekly_report.ai.clusterer import ActivityClusterer
            clusterer = ActivityClusterer(
                paths.LITELLM_CONFIG
            )
            if not clusterer.enabled:
                return []
            return clusterer.build_topics(activities)
        except Exception as e:
            print(f"Warning: topic clustering failed: {e}")
            return []

    def _issues_and_plan(self, activities, week_start, week_end):
        """STEP 2 '주요 이슈 & 리스크'·'다음 주 계획' (ai/issues.py). 실패·근거 없음이면 None (섹션 생략)."""
        if not week_start or not week_end:
            return None
        try:
            from weekly_report.ai.issues import IssuesAndPlan
            return IssuesAndPlan(self, paths.LITELLM_CONFIG).build(activities, week_start, week_end)
        except Exception as e:
            print(f"Warning: issues/plan failed: {e}")
            return None

    def _rag_sections(self, activities, week_start, week_end):
        """Settings의 주제별로 VectorDB 검색 → 근거 인용 요약 (rag.py).
        기간을 모르거나 LiteLLM 설정이 없거나 실패하면 빈 리스트 (섹션 생략)."""
        if not week_start or not week_end:
            return []
        try:
            from weekly_report.ai.rag import TopicQueryRAG
            rag = TopicQueryRAG(self, paths.LITELLM_CONFIG)
            return rag.build_sections(activities, week_start, week_end)
        except Exception as e:
            print(f"Warning: topic query (RAG) failed: {e}")
            return []

    def _expected_questions(self, activities, week_start, week_end, report_points):
        """STEP 3: 예상질문자별 예상 질문 & 답변 (ai/questions.py).
        예상질문자 미설정·LiteLLM 미설정·실패 시 빈 리스트 (섹션 생략)."""
        if not week_start or not week_end or not report_points:
            return []
        try:
            from weekly_report.ai.questions import ExpectedQuestions
            qa = ExpectedQuestions(self, paths.LITELLM_CONFIG)
            if not qa.enabled:
                return []
            return qa.build_sections(activities, week_start, week_end, report_points)
        except Exception as e:
            print(f"Warning: expected questions (STEP 3) failed: {e}")
            return []

    def _dedupe_meetings(self, activities):
        """같은 회의가 Outlook(COM)과 Teams(Graph) 양쪽으로 수집된 경우
        참석자 정보가 풍부한 Teams 쪽만 남긴다. 제목 기준으로 비교
        (타임스탬프는 Outlook=로컬, Graph=UTC라 단순비교 불가)."""
        teams_subjects = set()
        for activity in activities:
            if activity.get("source") != "teams" or activity.get("action") != "meeting":
                continue
            subject = self._meeting_subject(activity)
            if subject:
                teams_subjects.add(subject)
        if not teams_subjects:
            return activities
        return [
            a for a in activities
            if not (
                a.get("source") == "outlook"
                and a.get("action") == "meeting"
                and self._meeting_subject(a) in teams_subjects
            )
        ]

    def _meeting_subject(self, activity):
        """활동에서 회의 제목을 정규화해서 반환 (details.subject 우선,
        없으면 file_path에서 '회의:'/'Meeting:' 프리픽스 제거)."""
        details = activity.get("details")
        if isinstance(details, str):
            try:
                details = json.loads(details)
            except (TypeError, ValueError):
                details = None
        if isinstance(details, dict) and details.get("subject"):
            return str(details["subject"]).strip().lower()
        path = str(activity.get("file_path") or "")
        for prefix in ("Teams 회의:", "Teams Meeting:", "회의:", "Meeting:"):
            if path.startswith(prefix):
                return path[len(prefix):].strip().lower()
        return path.strip().lower() or None

    # 브라우저 URL 중 인증/세션 관련 URL (쿼리에 토큰·코드가 들어가 노출되면 안 됨)
    SENSITIVE_URL_MARKERS = (
        "/callback", "session_state", "access_token", "id_token", "code=",
    )

    def _is_noise_activity(self, activity):
        """오피스 임시/잠금 파일, DRM으로 깨진 파일명, 인증 콜백 URL 등
        사람이 봐도 의미 없거나 노출하면 안 되는 항목인지 판단"""
        if activity.get("source") == "browser":
            url = str(activity.get("file_path") or "").lower()
            return any(marker in url for marker in self.SENSITIVE_URL_MARKERS)
        if activity.get("source") != "filesystem":
            return False
        file_path = str(activity.get("file_path") or "")
        if ".git" in file_path.replace("\\", "/").split("/"):
            return True
        filename = os.path.basename(file_path)
        if self._HEX_TEMP_NAME_RE.match(filename):
            return True
        if "^" in filename:
            return True
        return any(
            fnmatch.fnmatch(filename, pattern)
            for pattern in self.NOISE_FILENAME_PATTERNS
        )

    def _parse_date(self, value):
        return datetime.strptime(value, "%Y-%m-%d")

    @staticmethod
    def _word_para(document, text, size=10.5, bold=False, italic=False,
                   indent=0.0, center=False, space_before=0):
        """글꼴 크기(pt)·굵기·들여쓰기(cm)를 직접 지정한 문단.
        자체 생성 템플릿에는 Heading 스타일이 없어서 크기를 run에 직접 준다."""
        paragraph = document.add_paragraph()
        if center:
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        fmt = paragraph.paragraph_format
        if indent:
            fmt.left_indent = Cm(indent)
        if space_before:
            fmt.space_before = Pt(space_before)
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        return paragraph

    def _ensure_word_template(self) -> str:
        template_path = Path(self.template_dir) / "blank_report_template.docx"
        if template_path.exists():
            try:
                Document(str(template_path))
                return str(template_path)
            except Exception:
                template_path.unlink(missing_ok=True)

        template_path.parent.mkdir(parents=True, exist_ok=True)
        self._write_minimal_docx(template_path)
        return str(template_path)

    def _write_minimal_docx(self, path: Path) -> None:
        """Build a minimal valid .docx from scratch.

        `Document()` (no args) relies on python-docx's own bundled
        `default.docx` template. On this machine that bundled file has been
        corrupted into an OLE2 container (signature `D0 CF 11 E0 ...`
        instead of the `PK` zip signature a .docx needs) — almost certainly
        by a corporate DRM/document-security tool that re-wraps .docx files
        saved to disk. Since we can't rely on that file being intact,
        generate a self-contained blank docx by hand instead.
        """
        content_types = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
            "</Types>"
        )
        package_rels = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>"
        )
        document_rels = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            "</Relationships>"
        )
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p/><w:sectPr/></w:body>"
            "</w:document>"
        )
        styles_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
            '<w:name w:val="Normal"/><w:qFormat/>'
            "</w:style>"
            "</w:styles>"
        )

        with ZipFile(path, "w", ZIP_DEFLATED) as docx_zip:
            docx_zip.writestr("[Content_Types].xml", content_types)
            docx_zip.writestr("_rels/.rels", package_rels)
            docx_zip.writestr("word/document.xml", document_xml)
            docx_zip.writestr("word/_rels/document.xml.rels", document_rels)
            docx_zip.writestr("word/styles.xml", styles_xml)


def main():
    generator = ReportGenerator()
    week_end = datetime.now().strftime("%Y-%m-%d")
    week_start = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
    result = generator.generate_and_save_weekly_report(
        week_start=week_start,
        week_end=week_end,
        formats=("markdown", "word"),
    )
    raw_file_paths = result.get("file_paths", {})
    if not isinstance(raw_file_paths, dict):
        raise RuntimeError("보고서 파일 경로 정보가 올바르지 않습니다.")
    file_paths = raw_file_paths

    print("주간 보고서 생성 완료")
    for report_format, file_path in file_paths.items():
        if file_path:
            print(f"- {report_format}: {file_path}")


if __name__ == "__main__":
    main()
