#!/usr/bin/env python3
"""
File System Watcher for Weekly Report Automation
Monitors file changes in specified directories and logs activity
"""

import os
import re
import sys
import time
import json
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from database import ActivityDatabase
from llm_summarizer import LLMSummarizer

# Notepad 등으로 계속 이어 적는 메모성 파일 확장자 - 실제로 무엇을 적었는지 캡처 대상
TEXT_CONTENT_EXTENSIONS = {'.txt', '.md'}
EXCEL_CONTENT_EXTENSIONS = {'.xlsx', '.xlsm'}
WORD_CONTENT_EXTENSIONS = {'.docx'}
PPT_CONTENT_EXTENSIONS = {'.pptx'}


class FileActivityHandler(FileSystemEventHandler):
    """Handler for file system events"""

    def __init__(self, log_file, exclude_patterns=None, db_path='data/activities.db'):
        self.log_file = log_file
        self.activities = []
        self.exclude_patterns = exclude_patterns or []
        # Always exclude log files and database files
        self.exclude_patterns.append('*.log')
        self.exclude_patterns.append('logs/*')
        self.exclude_patterns.append('*.db')
        self.exclude_patterns.append('*.db-journal')
        self.exclude_patterns.append('data/*')
        # 사람이 봐도 의미 없는 노이즈: 브라우저 다운로드 임시파일, git 내부 파일,
        # 오피스/OS가 저장할 때 쓰는 잠금·임시 파일
        self.exclude_patterns.append('*.crdownload')
        self.exclude_patterns.append('*.part')
        self.exclude_patterns.append('*.lock')
        self.exclude_patterns.append('.git')
        # report_generator.NOISE_FILENAME_PATTERNS와 동일 규칙 유지:
        # 오피스 잠금파일(~$*), 임시파일(*.tmp, *.cache), LibreOffice 잠금(.~lock.*)
        self.exclude_patterns.append('~$*')
        self.exclude_patterns.append('*.tmp')
        self.exclude_patterns.append('*.cache')
        self.exclude_patterns.append('.~lock.*')
        self.db_path = db_path
        self.db = None  # Will be created per thread
        self.summarizer = LLMSummarizer()
        # 파일별 마지막으로 읽은 텍스트 내용 (이번 세션 동안의 diff 계산용)
        self.last_seen_content = {}
        # 파일별 마지막으로 읽은 Excel 시트 스냅샷 ({sheet_name: [row_tuple, ...]})
        self.last_seen_excel = {}
        # 파일별 마지막으로 읽은 Word 문단 목록 / PPT 슬라이드 텍스트 목록
        self.last_seen_word = {}
        self.last_seen_ppt = {}

    # Office/OS가 원자적 저장(임시파일 쓰고 이름 바꿔치기) 중에 만드는
    # 8자리 16진수 임시 파일명 (예: "8C137000", "72EE75EC.tmp")
    _HEX_TEMP_NAME_RE = re.compile(r'^[0-9A-Fa-f]{8}(\.[A-Za-z0-9]+)?$')

    def should_exclude(self, file_path):
        """Check if file should be excluded from monitoring.

        패턴은 파일명(basename) 기준으로 매칭한다. 예전엔 전체 경로에
        fnmatch를 적용해서 "~*" 같은 패턴이 실제로는 전혀 매칭되지 않는
        버그가 있었음 (전체 경로가 "~"로 시작하는 경우는 없으니까).
        """
        import fnmatch
        filename = os.path.basename(file_path)

        if self._HEX_TEMP_NAME_RE.match(filename):
            return True

        # 사내 DRM/문서보안 솔루션이 파일을 재암호화하면서 "(DDr^7D3" 처럼
        # "^" 문자가 섞인 깨진 이름으로 바꿔놓는 경우가 있음 - 정상적인
        # 파일명에는 "^"가 거의 쓰이지 않으므로 노이즈로 간주
        if '^' in filename:
            return True

        for pattern in self.exclude_patterns:
            if fnmatch.fnmatch(filename, pattern) or pattern in file_path:
                return True
        return False
    
    def on_created(self, event):
        if not event.is_directory and not self.should_exclude(event.src_path):
            details = self._capture_file_content(event.src_path)
            self.log_activity('created', event.src_path, details)

    def on_modified(self, event):
        if not event.is_directory and not self.should_exclude(event.src_path):
            details = self._capture_file_content(event.src_path)
            self.log_activity('modified', event.src_path, details)

    def on_deleted(self, event):
        if not event.is_directory and not self.should_exclude(event.src_path):
            self.last_seen_content.pop(event.src_path, None)
            self.last_seen_excel.pop(event.src_path, None)
            self.last_seen_word.pop(event.src_path, None)
            self.last_seen_ppt.pop(event.src_path, None)
            self.log_activity('deleted', event.src_path)

    def on_moved(self, event):
        if not event.is_directory and not self.should_exclude(event.src_path):
            self.log_activity('moved', f"{event.src_path} -> {event.dest_path}")

    def _capture_file_content(self, file_path):
        """파일 확장자에 맞는 내용 캡처 방식으로 분기"""
        ext = Path(file_path).suffix.lower()
        if ext in EXCEL_CONTENT_EXTENSIONS:
            return self._capture_excel_content(file_path)
        if ext in WORD_CONTENT_EXTENSIONS:
            return self._capture_word_content(file_path)
        if ext in PPT_CONTENT_EXTENSIONS:
            return self._capture_ppt_content(file_path)
        return self._capture_text_content(file_path)

    def _diff_lines(self, old_lines, new_lines, limit=30):
        """new_lines 중 old_lines에 없던(새로 생기거나 바뀐) 항목만 순서대로 반환.
        Counter 기반이라 중복/순서변경도 적당히 잘 처리됨. 삭제된 내용은
        보고 대상이 아니므로 무시."""
        remaining = Counter(old_lines)
        diff = []
        for line in new_lines:
            if remaining[line] > 0:
                remaining[line] -= 1
            else:
                diff.append(line)
                if len(diff) >= limit:
                    break
        return diff

    def _finalize_capture(self, added_lines, joiner="\n"):
        """diff된 줄 목록을 요약 details(dict)로 변환. 없으면 None."""
        if not added_lines:
            return None
        added = joiner.join(added_lines)
        summary = self.summarizer.summarize(added, template="file") or added[:300]
        return {'content_added': added[:2000], 'summary': summary}

    def _drm_blocked_details(self):
        """사내 DRM/문서보안 솔루션이 파일을 OLE2로 재암호화해서 zip으로
        못 여는 경우: 변경이 있었다는 사실만이라도 보고서에 남기기 위한
        fallback details. 파일이 다른 프로그램에 열려 있어 일시적으로
        잠긴 경우(PermissionError 등)는 이 경로를 타지 않음 - 그런 경우는
        다음 저장 시 다시 시도하면 되므로 조용히 None을 반환하는 게 맞음."""
        return {
            'content_added': None,
            'summary': '(내용 변경 감지됨 - 사내 보안 정책으로 암호화되어 상세 내용 확인 불가)',
        }

    # --- Windows Office COM 폴백 (DRM/OLE2 파일 읽기) ---

    def _open_office_com(self, prog):
        """Windows Office COM 앱 인스턴스 반환. 비Windows/미설치 시 None."""
        if sys.platform != "win32":
            return None
        try:
            import win32com.client
            app = win32com.client.Dispatch(f"{prog}.Application")
            for attr, val in (("Visible", False), ("DisplayAlerts", False)):
                try:
                    setattr(app, attr, val)
                except Exception:
                    pass
            return app
        except Exception:
            return None

    def _excel_snapshot_via_com(self, file_path):
        """OLE2로 재암호화된 xlsx를 Excel COM으로 열어 openpyxl과 동일한
        스냅샷 dict 반환. Excel이 자체 권한으로 복호화하므로 라벨 해제
        없이 내용이 읽힘. 실패 시 None."""
        app = self._open_office_com("Excel")
        if app is None:
            return None
        try:
            wb = app.Workbooks.Open(file_path, ReadOnly=True)
            snapshot = {}
            for ws in wb.Worksheets:
                rows = []
                for row in ws.UsedRange.Rows:
                    vals = tuple(c.Value for c in row.Cells)
                    if any(v is not None and str(v).strip() for v in vals):
                        rows.append(vals)
                snapshot[ws.Name] = rows
            wb.Close(False)
            return snapshot
        except Exception:
            return None
        finally:
            try:
                app.Quit()
            except Exception:
                pass

    def _word_paragraphs_via_com(self, file_path):
        """OLE2 docx를 Word COM으로 열어 문단 텍스트 리스트 반환."""
        app = self._open_office_com("Word")
        if app is None:
            return None
        try:
            doc = app.Documents.Open(file_path, ReadOnly=True)
            paras = [p.Range.Text.strip() for p in doc.Paragraphs
                     if p.Range.Text.strip()]
            doc.Close(False)
            return paras
        except Exception:
            return None
        finally:
            try:
                app.Quit()
            except Exception:
                pass

    def _ppt_lines_via_com(self, file_path):
        """OLE2 pptx를 PowerPoint COM으로 열어 슬라이드 텍스트 리스트 반환."""
        app = self._open_office_com("PowerPoint")
        if app is None:
            return None
        try:
            pres = app.Presentations.Open(file_path, ReadOnly=True,
                                          WithWindow=False)
            lines = []
            for slide_number, slide in enumerate(pres.Slides, start=1):
                texts = []
                for shape in slide.Shapes:
                    try:
                        if shape.HasTextFrame and shape.TextFrame.HasText:
                            text = shape.TextFrame.TextRange.Text.strip()
                            if text:
                                texts.append(text)
                    except Exception:
                        continue
                if texts:
                    lines.append(f"[슬라이드 {slide_number}] " + " / ".join(texts))
            pres.Close()
            return lines
        except Exception:
            return None
        finally:
            try:
                app.Quit()
            except Exception:
                pass

    def _capture_word_content(self, file_path):
        """Word(.docx) 파일에서 새로 추가/변경된 문단을 캡처.

        사내 DRM/문서보안 솔루션이 저장 직후 .docx를 OLE2 컨테이너로
        재암호화하는 경우가 있어(이 프로젝트에서 이미 확인된 이슈),
        파서 실패 시 Windows에서는 Word COM으로 한 번 더 읽고,
        그것도 실패하면 "보안 문서" 폴백 details를 반환한다.
        """
        try:
            from docx import Document
        except ImportError:
            return None
        try:
            doc = Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs
                          if p.text.strip()]
        except zipfile.BadZipFile:
            paragraphs = self._word_paragraphs_via_com(file_path)
            if paragraphs is None:
                return self._drm_blocked_details()
        except Exception:
            return None

        previous = self.last_seen_word.get(file_path, [])
        self.last_seen_word[file_path] = paragraphs

        return self._finalize_capture(self._diff_lines(previous, paragraphs))

    def _capture_ppt_content(self, file_path):
        """PowerPoint(.pptx) 파일에서 새로 추가/변경된 슬라이드 텍스트를 캡처."""
        try:
            from pptx import Presentation
        except ImportError:
            return None
        try:
            presentation = Presentation(file_path)
            lines = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        text = shape.text_frame.text.strip()
                        if text:
                            texts.append(text)
                if texts:
                    lines.append(f"[슬라이드 {slide_number}] " + " / ".join(texts))
        except zipfile.BadZipFile:
            lines = self._ppt_lines_via_com(file_path)
            if lines is None:
                return self._drm_blocked_details()
        except Exception:
            return None

        previous = self.last_seen_ppt.get(file_path, [])
        self.last_seen_ppt[file_path] = lines

        return self._finalize_capture(self._diff_lines(previous, lines, limit=20))

    def _capture_excel_content(self, file_path):
        """Excel(.xlsx) 파일에서 새로 추가/변경된 행을 캡처해서
        details(added/summary)로 반환. 읽기 실패 시 None."""
        try:
            import openpyxl
        except ImportError:
            return None

        snapshot = None
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            max_rows, max_cols = 500, 50
            snapshot = {}
            try:
                for sheet in workbook.worksheets:
                    rows = []
                    for row in sheet.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True):
                        if any(cell is not None and str(cell).strip() for cell in row):
                            rows.append(tuple(row))
                    snapshot[sheet.title] = rows
            finally:
                workbook.close()
        except zipfile.BadZipFile:
            snapshot = self._excel_snapshot_via_com(file_path)
            if snapshot is None:
                return self._drm_blocked_details()
        except Exception:
            return None

        previous = self.last_seen_excel.get(file_path, {})
        self.last_seen_excel[file_path] = snapshot

        changed_lines = []
        for sheet_name, rows in snapshot.items():
            old_rows = previous.get(sheet_name, [])
            for idx, row in enumerate(rows):
                old_row = old_rows[idx] if idx < len(old_rows) else None
                if row == old_row:
                    continue
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    changed_lines.append(f"[{sheet_name}] " + " | ".join(cells))

        return self._finalize_capture(changed_lines[:20])

    def _capture_text_content(self, file_path):
        """Notepad류 텍스트 파일(.txt/.md)에 새로 적힌 내용을 캡처해서
        details(added/summary)로 반환. 대상이 아니거나 읽기 실패 시 None."""
        ext = Path(file_path).suffix.lower()
        if ext not in TEXT_CONTENT_EXTENSIONS:
            return None

        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        except (OSError, UnicodeError):
            return None

        previous = self.last_seen_content.get(file_path, '')
        self.last_seen_content[file_path] = content

        if content.startswith(previous):
            added = content[len(previous):].strip()
        else:
            # 중간 수정/재정렬 등으로 단순 이어쓰기가 아니면 전체 내용을 새 내용으로 취급
            added = content.strip()

        if not added:
            return None

        summary = self.summarizer.summarize(added, template="file") or added[:300]
        return {'content_added': added[:2000], 'summary': summary}

    def log_activity(self, action, file_path, details=None):
        """Log file activity to file, memory, and database"""
        timestamp = datetime.now().isoformat()
        activity = {
            'timestamp': timestamp,
            'action': action,
            'file_path': file_path,
            'file_type': self.get_file_type(file_path)
        }
        if details:
            activity['details'] = json.dumps(details, ensure_ascii=False)

        self.activities.append(activity)

        # Write to log file
        log_entry = f"{timestamp} | {action.upper():10} | {file_path}\n"
        with open(self.log_file, 'a', encoding='utf-8') as f:
            f.write(log_entry)

        # Save to database (create new connection per thread)
        try:
            with ActivityDatabase(self.db_path) as db:
                db.add_activity(activity)
        except Exception as e:
            print(f"Error saving to database: {e}")

        print(f"[{timestamp}] {action.upper()}: {file_path}")
    
    def get_file_type(self, file_path):
        """Determine file type based on extension"""
        ext = Path(file_path).suffix.lower()
        type_map = {
            '.xlsx': 'excel',
            '.xls': 'excel',
            '.pptx': 'powerpoint',
            '.ppt': 'powerpoint',
            '.docx': 'word',
            '.doc': 'word',
            '.txt': 'text',
            '.md': 'markdown',
            '.py': 'python',
            '.js': 'javascript',
            '.json': 'json',
            '.csv': 'csv'
        }
        return type_map.get(ext, 'unknown')


class FileWatcher:
    """Main file watcher class"""
    
    def __init__(self, watch_paths, log_file='logs/file_activity.log', exclude_patterns=None, db_path='data/activities.db'):
        self.watch_paths = watch_paths
        self.log_file = log_file
        self.exclude_patterns = exclude_patterns or []
        self.db_path = db_path
        self.observer = Observer()
        self.handler = None
        self._running = False
        
        # Ensure log directory exists
        os.makedirs(os.path.dirname(log_file), exist_ok=True)
    
    def start(self):
        """Start watching directories"""
        self.handler = FileActivityHandler(self.log_file, self.exclude_patterns, self.db_path)
        
        for path in self.watch_paths:
            if os.path.exists(path):
                self.observer.schedule(self.handler, path, recursive=True)
                print(f"Watching: {path}")
            else:
                print(f"Warning: Path does not exist: {path}")
        
        self.observer.start()
        self._running = True
        print(f"File watcher started. Monitoring {len(self.watch_paths)} directory/directories.")
        print("Press Ctrl+C to stop...")

        try:
            while self._running:
                time.sleep(1)
        except KeyboardInterrupt:
            self.stop()
    
    def stop(self):
        """Stop watching directories"""
        self._running = False
        self.observer.stop()
        self.observer.join()
        print("\nFile watcher stopped.")
        
        # Save activities to JSON
        self.save_activities()
    
    def save_activities(self):
        """Save collected activities to JSON file"""
        if self.handler and self.handler.activities:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            json_file = f"data/activities_{timestamp}.json"
            
            os.makedirs(os.path.dirname(json_file), exist_ok=True)
            
            with open(json_file, 'w', encoding='utf-8') as f:
                json.dump(self.handler.activities, f, indent=2, ensure_ascii=False)
            
            print(f"Activities saved to: {json_file}")


def load_config(config_file='config/watch_config.json'):
    """Load watch paths from config file"""
    if os.path.exists(config_file):
        with open(config_file, 'r', encoding='utf-8') as f:
            config = json.load(f)
            return config.get('watch_paths', []), config.get('exclude_patterns', [])
    return [], []


def main():
    """Main entry point"""
    # Load config or use default paths
    watch_paths, exclude_patterns = load_config()
    
    if not watch_paths:
        # Default paths for testing
        print("No config found. Using default paths...")
        watch_paths = [
            os.path.expanduser("~/Documents"),
            os.path.expanduser("~/Desktop")
        ]
    
    # Create and start watcher
    watcher = FileWatcher(watch_paths, exclude_patterns=exclude_patterns)
    watcher.start()


if __name__ == "__main__":
    main()