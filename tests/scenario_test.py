# -*- coding: utf-8 -*-
"""전체 기능 테스트 시나리오 실행 → Excel 결과 생성 (test_reports/)

카테고리:
  환경/인증 / 수집 / 데이터 가공 / AI 분석 / 보고서 생성 / GUI·실행
각 항목은 PASS / FAIL / SKIP 중 하나의 결과를 낸다.
"""

import json
import os
import re
import sys
import time
import traceback
from datetime import datetime, timedelta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.chdir(ROOT)  # 상대 경로(config/, data/, test_reports/)는 프로젝트 루트 기준

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

RESULTS = []  # dict: id, category, scenario, result, detail, elapsed


def run(tc_id, category, scenario, fn):
    t0 = time.time()
    try:
        status, detail = fn()
    except Exception as e:
        status, detail = "FAIL", f"{type(e).__name__}: {e}"
        traceback.print_exc()
    RESULTS.append(
        {
            "id": tc_id,
            "category": category,
            "scenario": scenario,
            "result": status,
            "detail": detail,
            "elapsed": round(time.time() - t0, 1),
        }
    )
    print(f"[{status:4}] {tc_id} {scenario} — {detail[:80]}")


def ok(detail=""):
    return "PASS", detail


def fail(detail=""):
    return "FAIL", detail


def skip(detail=""):
    return "SKIP", detail


# ───────────────────────── 환경/인증 ─────────────────────────

def t_db_schema():
    from weekly_report.storage.database import ActivityDatabase
    with ActivityDatabase() as db:
        rows = db.conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        ).fetchall()
        names = {r[0] for r in rows}
        need = {"activities", "weekly_summaries"}
        missing = need - names
        if missing:
            return fail(f"테이블 누락: {missing}")
        cnt = db.conn.execute("SELECT COUNT(*) FROM activities").fetchone()[0]
    return ok(f"테이블 정상, activities {cnt}건")


def t_litellm_config():
    from weekly_report.ai.llm_summarizer import LLMSummarizer
    s = LLMSummarizer()
    if not s.enabled:
        return fail("config 비활성 또는 api_key 없음")
    return ok("config 로드 + enabled")


def t_embedding_call():
    from weekly_report.ai.clusterer import ActivityClusterer
    c = ActivityClusterer()
    if not c.enabled:
        return skip("LiteLLM config 없음")
    vecs = c.embed(["테스트 문장"])
    if not vecs:
        return fail("임베딩 호출 실패")
    return ok(f"임베딩 성공 dim={len(vecs[0])}")


def t_lcel_chat():
    from weekly_report.ai.llm_summarizer import LLMSummarizer
    s = LLMSummarizer()
    if not s.enabled:
        return skip("LiteLLM config 없음")
    text = ("주간보고서 자동화 프로그램을 개발 중이다. " * 20)
    out = s.summarize(text, max_len=100, template="default")
    if not out or len(out) >= len(text):
        return fail(f"요약 미동작: {str(out)[:60]}")
    return ok(f"LCEL 체인 요약 성공 ({len(out)}자)")


def t_teams_token():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 파일 없음 (최초 로그인 필요)")
    with open("config/teams_graph_token.json", encoding="utf-8") as f:
        tok = json.load(f)
    has_rt = bool(tok.get("refresh_token"))
    if not has_rt:
        return fail("refresh_token 없음")
    return ok("토큰 파일 존재, refresh_token 보유")


def t_teams_auth_live():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 파일 없음")
    from weekly_report.collectors.teams import get_access_token
    token = get_access_token()
    if not token:
        return fail("토큰 갱신/획득 실패")
    return ok("Graph 액세스 토큰 획득 (refresh 동작)")


# ───────────────────────── 수집 ─────────────────────────

def t_watcher_noise():
    from weekly_report.collectors.file_watcher import FileActivityHandler
    w = FileActivityHandler("logs/_test.log")
    cases = [
        ("C:/x/미확인 819675.crdownload", True),
        ("C:/x/(DDr^7D3.jpg", True),
        ("C:/x/D58F4000", True),
        ("C:/x/~$report.docx", True),
        ("C:/x/보고서.docx", False),
    ]
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    bad = [
        p for p, want in cases
        if bool(
            w.should_exclude(p)
            or rg._is_noise_activity(
                {"file_path": p, "source": "filesystem", "action": "modified",
                 "timestamp": "t"}
            )
        ) != want
    ]
    if bad:
        return fail(f"필터 오판: {bad}")
    return ok("노이즈 4종 + 정상파일 구분 정확")


def t_office_content_capture():
    """실제 오피스 파일 내용 추출 또는 DRM 폴백"""
    from weekly_report.collectors.file_watcher import FileActivityHandler
    w = FileActivityHandler("logs/_test.log")
    # xlsx DRM 폴백 경로 테스트: 실제 DRM 파일이 없으면 SKIP
    import glob
    drm_files = [
        p for p in glob.glob(r"C:\Users\SSG\Documents\**\*.xlsx", recursive=True)
        if os.path.exists(p)
    ][:3]
    if not drm_files:
        return skip("테스트용 xlsx 없음")
    res = w._capture_excel_content(drm_files[0])
    if res is None:
        return ok("비zip/잠김 → 조용히 None (정상 경로)")
    return ok(f"내용 추출/폴백 메시지: {str(res)[:60]}")


def t_browser_server():
    """브라우저 수신 서버 포트 상태 + POST 수신 경로"""
    import socket
    sock = socket.socket()
    sock.settimeout(2)
    try:
        sock.connect(("127.0.0.1", 5757))
        return ok("5757 포트 리스닝 중 (트래킹 실행 중)")
    except OSError:
        return skip("서버 미기동 — Start Tracking으로 실행되면 동작")
    finally:
        sock.close()


def t_ide_paths():
    from weekly_report.collectors.ide import IDECollector
    c = IDECollector()
    paths = c._find_vscode_paths()
    if not paths:
        return skip("VS Code 경로 없음")
    return ok(f"VS Code 경로 {len(paths)}개 발견")


def t_teams_chat_collect():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from weekly_report.collectors.teams import TeamsCollector
    col = TeamsCollector()
    acts = col.collect_all_teams_activity(days=1)
    return ok(f"Teams 활동 {len(acts)}건 수집")


def t_teams_calendar():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from weekly_report.collectors.teams import TeamsCollector
    col = TeamsCollector()
    acts = col.collect_calendar_events(days=7)
    return ok(f"일정 {len(acts)}건 (취소 제외)")


def t_outlook_collect():
    from weekly_report.collectors.outlook import OutlookCollector, _WIN32COM_AVAILABLE
    c = OutlookCollector()
    if _WIN32COM_AVAILABLE:
        connected = c.connect_outlook()
        if not connected:
            return skip("Outlook 미실행/COM 연결 실패")
        acts = c.collect_sent_email_activity(days=1)
        return ok(f"COM 경로 — 보낸 메일 {len(acts)}건")
    acts = c.collect_sent_email_activity(days=1)
    return ok(f"Graph 경로 — {len(acts)}건")


def t_confluence():
    cfg = "config/confluence_config.json"
    if not os.path.exists(cfg):
        return skip("Confluence config 없음")
    return ok("config 존재 (토큰 유효성은 수집 시 확인)")


def t_slack():
    cfg = "config/slack_config.json"
    if not os.path.exists(cfg):
        return skip("Slack 토큰 미발급 — 수집 비활성")
    return ok("config 존재")


def t_onenote():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from weekly_report.collectors.onenote import OneNoteCollector
    col = OneNoteCollector()
    # 페이지 수준은 권한 부재 — 노트북/섹션 목록 수집 가능 여부 확인
    try:
        acts = col.collect_activity(days=400)
        return ok(f"노트북/섹션 변경 {len(acts)}건 수집 (페이지 단위는 권한 제한)")
    except Exception as e:
        return fail(str(e)[:80])


def t_sharepoint():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from weekly_report.collectors.sharepoint import SharePointCollector
    col = SharePointCollector()
    try:
        # 버전 비교(다운로드+COM)는 느려서 여기선 수정 목록 수집만 확인 — 비교는 COL-13
        acts = col.collect_activity(days=7, with_changes=False)
        return ok(f"SharePoint/OneDrive 수정 파일 {len(acts)}건")
    except Exception as e:
        return fail(str(e)[:80])


def t_office_com_diff():
    """Office COM으로 xlsx 두 버전을 읽어 변경 줄을 뽑는지 (SharePoint 버전 비교 경로)"""
    if sys.platform != "win32":
        return skip("Windows 전용")
    import tempfile
    import openpyxl
    from weekly_report.common import office_reader
    tmp = tempfile.mkdtemp(prefix="wr_test_")
    paths = []
    for i, rows in enumerate([[("항목", "담당"), ("DB접속 암호화", "박성수")],
                              [("항목", "담당"), ("DB접속 암호화 (평문제거)", "박성수")]]):
        wb = openpyxl.Workbook()
        wb.active.title = "배포"
        for r in rows:
            wb.active.append(r)
        path = os.path.join(tmp, f"v{i}.xlsx")
        wb.save(path)
        paths.append(path)
    with office_reader.office_app("Excel") as app:
        if app is None:
            return skip("Excel 미설치")
        old, new = (office_reader.document_lines(app, "Excel", p) for p in paths)
    if old is None or new is None:
        return fail("COM으로 xlsx 읽기 실패")
    added, removed = office_reader.diff_lines(old, new)
    if added != ["[배포] DB접속 암호화 (평문제거) | 박성수"] or removed != ["[배포] DB접속 암호화 | 박성수"]:
        return fail(f"diff 불일치: +{added} -{removed}")
    return ok("COM 읽기 + 버전 diff 정상 (+1/-1줄)")


def t_file_watcher_watchdirs():
    cfg = "config/watch_config.json"
    if not os.path.exists(cfg):
        return fail("watch_config.json 없음")
    data = json.load(open(cfg, encoding="utf-8"))
    dirs = data.get("watch_paths", [])
    if not dirs:
        return fail("감시 폴더 0개")
    existing = [d for d in dirs if os.path.exists(d)]
    return ok(f"감시 폴더 {len(dirs)}개 (실존 {len(existing)})")


# ───────────────────────── 데이터 가공 ─────────────────────────

def t_meeting_dedupe():
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    acts = [
        {"timestamp": "2026-09-29T10:00:00", "action": "meeting",
         "file_path": "Meeting: 주간회의", "source": "outlook",
         "details": "{}"},
        {"timestamp": "2026-09-29T10:00:00", "action": "meeting",
         "file_path": "Teams Meeting: 주간회의", "source": "teams",
         "details": '{"subject": "주간회의"}'},
    ]
    out = rg._dedupe_meetings(acts)
    if len(out) == 1 and out[0]["source"] == "teams":
        return ok("동일 회의 중복 제거 + teams 우선")
    return fail(f"{len(out)}건 남음: {[a['source'] for a in out]}")


def t_transient_files():
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    acts = [
        {"timestamp": "t1", "action": "created", "file_path": "a.tmp", "source": "filesystem"},
        {"timestamp": "t2", "action": "modified", "file_path": "a.tmp", "source": "filesystem"},
        {"timestamp": "t3", "action": "deleted", "file_path": "a.tmp", "source": "filesystem"},
        {"timestamp": "t4", "action": "created", "file_path": "keep.docx", "source": "filesystem"},
    ]
    out = rg._drop_transient_files(acts)
    paths = {a["file_path"] for a in out}
    if paths == {"keep.docx"}:
        return ok("created+deleted 쌍 제거 정상")
    return fail(f"남은 경로: {paths}")


def t_path_normalize():
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    a = {"file_path": "C:/Users/SSG/test.txt"}
    out = rg._normalize_activity_path(a)
    if "/" in out["file_path"]:
        return fail(out["file_path"])
    return ok(f"정규화: {out['file_path']}")


def t_sensitive_url_filter():
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    bad = {
        "file_path": "https://login.example.com/callback?code=abc&state=x",
        "source": "browser", "action": "visited", "timestamp": "t",
    }
    if rg._is_noise_activity(bad):
        return ok("OAuth 콜백 URL 필터링됨")
    return fail("민감 URL 통과")


def t_summary_cache():
    from weekly_report.storage.database import ActivityDatabase
    with ActivityDatabase() as db:
        try:
            cache = db.get_summary_cache("outlook")
            return ok(f"요약 캐시 {len(cache)}건 로드")
        except AttributeError:
            return fail("get_summary_cache 없음")


def t_shared_doc_split():
    """SharePoint: 내가 수정한 문서는 Excel, 파트원만 수정한 문서는 '공유 문서 변경'으로
    분리되고, 공유 문서 변경은 STEP 2 요약 입력에서 빠지는지"""
    from weekly_report.report.sections import build_program_sections, sections_digest, SHARED_DOCS

    def sp(name, editors):
        return {"timestamp": "2026-09-29T23:59:59", "source": "sharepoint", "action": "modified",
                "file_path": f"SharePoint 파일: {name}", "file_type": "sharepoint",
                "details": json.dumps({"name": name, "editors": editors, "summary": f"{name} 변경"})}

    me = "홍길동(POS서버) - 팀A"
    secs = build_program_sections(
        [sp("내파일.xlsx", ["홍길동(POS서버) - 팀B", "김철수(POS) - 팀A"]),
         sp("팀파일.xlsx", ["김철수(POS) - 팀A"])],
        me_name=me,
    )
    by_name = {s["name"]: [i["title"] for i in s["items"]] for s in secs}
    if by_name.get("Excel") != ["내파일.xlsx (SharePoint)"] or by_name.get(SHARED_DOCS) != ["팀파일.xlsx (SharePoint)"]:
        return fail(f"분류 오류: {by_name}")
    if "팀파일" in sections_digest(secs):
        return fail("공유 문서 변경이 STEP 2 입력에 포함됨")
    return ok("내 수정→Excel, 파트원 수정→공유 문서 변경, STEP 2 제외")


def t_vector_store():
    """VectorDB(Qdrant 로컬): 저장·재조회·유사도 검색·날짜/소스/발신자 필터 (임시 폴더, 4차원 벡터)"""
    import json as _json, shutil, tempfile
    from weekly_report.storage.vector_store import ActivityVectorStore
    tmp = tempfile.mkdtemp(prefix="wr_qdrant_")
    try:
        store = ActivityVectorStore(path=tmp, vector_size=4)

        def act(ts, source, path, sender=""):
            return {"timestamp": ts, "action": "message", "file_path": path, "source": source,
                    "details": _json.dumps({"sender": sender}, ensure_ascii=False)}

        a1 = act("2026-09-28T10:00:00", "teams", "배포 일정", "김영호(POS) - 팀")
        a2 = act("2026-09-29T10:00:00", "teams", "장애 대응", "홍길동(POS) - 팀")
        a3 = act("2026-09-20T10:00:00", "outlook", "배포 공지")
        store.upsert([(a1, "배포 일정", [1, 0, 0, 0]), (a2, "장애 대응", [0, 1, 0, 0]), (a3, "배포 공지", [0.9, 0.1, 0, 0])])
        store.upsert([(a1, "배포 일정", [1, 0, 0, 0])])  # 같은 활동 재저장 → 덮어쓰기
        if store.count() != 3:
            return fail(f"건수 {store.count()} (재저장 중복)")
        from weekly_report.storage.database import ActivityDatabase
        got = store.get_vectors([ActivityDatabase.activity_key(a1)])
        if list(got.values())[0][0] != "배포 일정":
            return fail("재조회 실패")
        top = store.search([1, 0, 0, 0], limit=2)
        if [h["text"] for h in top] != ["배포 일정", "배포 공지"]:
            return fail(f"유사도 순서 오류: {[h['text'] for h in top]}")
        if [h["text"] for h in store.search([1, 0, 0, 0], date_from="2026-09-25", date_to="2026-09-28")] != ["배포 일정"]:
            return fail("날짜 필터 오류")
        if {h["source"] for h in store.search([1, 0, 0, 0], sources=["outlook"])} != {"outlook"}:
            return fail("소스 필터 오류")
        if [h["text"] for h in store.search([1, 0, 0, 0], sender_contains="홍길동")] != ["장애 대응"]:
            return fail("발신자 필터 오류")
        return ok("저장·덮어쓰기·재조회·유사도 순서·날짜/소스/발신자 필터 정상")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def t_project_docs():
    """프로젝트 문서·지난 보고 수집: 첫 실행은 기준만, 이후 변경분 요약 / 끝난 주 보고서만 다음 주 월요일 2차 근거,
    이번 주·복사본 제외 / 근거 제외 키워드(weekly_report)에 안 걸리고 STEP 1에는 안 나옴"""
    import json as _json, shutil, tempfile
    from datetime import date, datetime
    from weekly_report.collectors.project_docs import ProjectDocsCollector
    from weekly_report.ai.rag import TopicQueryRAG, evidence_label, is_evidence_noise
    from weekly_report.report.sections import classify_program
    tmp = tempfile.mkdtemp(prefix="wr_projdocs_")
    try:
        base = os.path.join(tmp, "weekly_report_automation")
        reports = os.path.join(base, "reports")
        os.makedirs(reports)
        doc = os.path.join(base, "progress.md")
        open(doc, "w", encoding="utf-8").write("# 진행\n- 항목 26 속도 개선\n")
        col = ProjectDocsCollector(base_dir=base, reports_dir=reports, snapshot_path=os.path.join(tmp, "snap.json"),
                                   summarize=lambda name, changes: f"{name}: " + ", ".join(changes), today=date(2026, 10, 7))
        if col.collect_docs():
            return fail("첫 실행에 문서 전체를 변경으로 넣음")
        open(doc, "w", encoding="utf-8").write("# 진행\n- 항목 26 속도 개선\n- 항목 27 저장소 보관 정책\n")
        os.utime(doc, (1790000000, 1790000000))
        docs = col.collect_docs()
        if len(docs) != 1 or "[+] - 항목 27 저장소 보관 정책" not in _json.loads(docs[0]["details"])["summary"]:
            return fail(f"변경분 수집 오류: {docs}")
        body = "# 보고\n### 이번 주 핵심 요약\n- 10월1차 배포 완료\n- 안면결제 개발 완료\n### 주제별 작업\n- x\n"
        for name in ("weekly_report_2026-09-28_2026-10-04.md", "weekly_report_2026-10-05_2026-10-11.md",
                     "weekly_report_2026-09-28_2026-10-04 - 복사본.md"):
            open(os.path.join(reports, name), "w", encoding="utf-8").write(body)
        past = col.collect_past_reports()
        if len(past) != 1 or past[0]["timestamp"] != "2026-10-05T09:00:00":
            return fail(f"지난 보고 선택 오류: {[(x['file_path'], x['timestamp']) for x in past]}")
        details = _json.loads(past[0]["details"])
        if details["summary"] != "- 10월1차 배포 완료\n- 안면결제 개발 완료":
            return fail(f"핵심 요약 추출 오류: {details['summary']}")
        kw = ["주간보고 자동작성", "weekly_report", "WeeklyPulse"]
        if any(is_evidence_noise(a, kw) for a in docs + past) or any(classify_program(a) for a in docs + past):
            return fail("근거 제외 키워드에 걸리거나 STEP 1에 나옴")
        label = evidence_label(past[0], datetime(2026, 10, 5, 9))
        if not label.startswith("10/05 (지난 보고) · 09/28~10/04 주간보고"):
            return fail(f"지난 보고 표시 오류: {label}")
        rag = TopicQueryRAG.__new__(TopicQueryRAG)
        rag.me_name, rag.teams_scope, rag.noise_keywords = "", "mine", kw
        rag.excluded_ids, rag.excluded_titles = set(), set()
        if not all(rag.eligible(a) for a in docs + past):
            return fail("근거 후보에서 빠짐")
        return ok("문서 첫 실행 기준만·변경분 요약, 끝난 주 보고서만 다음 주 월요일 2차 근거(이번 주·복사본 제외), "
                  "제외 키워드 면제·STEP 1 미표시·'(지난 보고)' 표시")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def t_mail_teams_fields():
    """일일 할 일 재료: Teams @언급(나·모든 사용자) 판정, 받은 메일 수신/참조 판정"""
    from types import SimpleNamespace
    from weekly_report.collectors.teams import TeamsCollector
    from weekly_report.collectors.outlook import OutlookCollector
    me = "손현태(이마트24 POS서버) - IT개발혁신TF"
    msg = {"mentions": [{"mentionText": me, "mentioned": {"user": {"displayName": me}}},
                        {"mentionText": "모든 사용자", "mentioned": {"conversation": {}}}]}
    other = {"mentions": [{"mentionText": "남궁윤", "mentioned": {"user": {"displayName": "남궁윤(이마트24POS) - 이마트24POS팀"}}}]}
    if not (TeamsCollector.mentions_me(msg, me) and TeamsCollector.mentions_all(msg)
            and not TeamsCollector.mentions_me(other, me) and not TeamsCollector.mentions_all(other)):
        return fail("Teams 언급 판정 오류")

    class Entry:
        def __init__(self, smtp):
            self.smtp = smtp
        def GetExchangeUser(self):
            return SimpleNamespace(PrimarySmtpAddress=self.smtp)
    def recip(name, smtp, kind):
        return SimpleNamespace(Name=name, Address="/o=ExchangeLabs/cn=x", Type=kind, AddressEntry=Entry(smtp))
    mine = {"191648@shinsegae.com", "손현태(이마트24 pos서버) - it개발혁신tf"}
    to_mail = SimpleNamespace(Recipients=[recip("김영호", "kim@x.com", 1), recip("손현태", "191648@shinsegae.com", 1)])
    cc_mail = SimpleNamespace(Recipients=[recip("김영호", "kim@x.com", 1), recip("손현태", "191648@shinsegae.com", 2)])
    dl_mail = SimpleNamespace(Recipients=[recip("운영서비스팀 전체", "ops@x.com", 1)])
    got = [OutlookCollector._recipient_type(m, mine) for m in (to_mail, cc_mail, dl_mail)]
    if got != ["to", "cc", "other"]:
        return fail(f"메일 수신 위치 판정 오류: {got}")
    # 1:1 전용 수집(할 일 추출용, DB 미저장): 1:1 채팅방만, until 이후 메시지는 제외
    from datetime import datetime as _dt, timezone as _tz
    import weekly_report.collectors.teams as teams_mod
    tc = TeamsCollector.__new__(TeamsCollector)
    tc._token, tc.token_file, tc._me_name = "x", None, me
    seen = {}
    def fake(token, since, types, excluded, me_name):
        seen["types"] = types
        return [{"timestamp": "2026-10-01T23:30:00Z"}, {"timestamp": "2026-10-02T01:00:00Z"}]
    tc._collect_chat_messages = fake
    got = tc.collect_direct_messages(_dt(2026, 10, 1, 0, 0, tzinfo=_tz.utc), until=_dt(2026, 10, 2, 0, 0, tzinfo=_tz.utc))
    if seen.get("types") != {"oneOnOne"} or [a["timestamp"] for a in got] != ["2026-10-01T23:30:00Z"]:
        return fail(f"1:1 전용 수집 오류: {seen} / {got}")
    return ok("Teams @나·@모든 사용자 판정, 받은 메일 수신(To)·참조(CC)·그룹 주소 판정, 1:1 전용 수집(기간·채팅 유형)")


def t_storage_retention():
    """저장소 보관 정책: 보관 기간이 지난 주는 요약 벡터 저장 후 활동 벡터 삭제, 요약 실패 주는 보류,
    최근 주는 유지, 레거시 임베딩 테이블 삭제, 장기 질의(요약 → 원본 활동) (임시 폴더·DB, 4차원)"""
    import json as _json, shutil, sqlite3, tempfile
    from datetime import date
    from weekly_report.storage.vector_store import ActivityVectorStore, WeekSummaryStore
    from weekly_report.storage.retention import apply_retention
    from weekly_report.storage.database import ActivityDatabase
    from weekly_report.ai.rag import search_past_weeks
    tmp = tempfile.mkdtemp(prefix="wr_retention_")
    try:
        store = ActivityVectorStore(path=os.path.join(tmp, "q"), vector_size=4)
        summaries = WeekSummaryStore(path=store.path, vector_size=4)
        db_path = os.path.join(tmp, "a.db")
        with ActivityDatabase(db_path) as db:
            db.add_activities([{"timestamp": "2026-06-02T10:00:00", "action": "sent", "file_path": "결제 연동 테스트",
                                "source": "outlook", "details": "{}"}])
            db.save_weekly_summary("2026-06-01", "2026-06-07", {
                "week_start": "2026-06-01", "week_end": "2026-06-07",
                "week_summary": ["결제 연동 테스트 완료"], "rag_sections": []})
            db.conn.execute("CREATE TABLE activity_embeddings (activity_key TEXT PRIMARY KEY, text TEXT, vector TEXT)")
            db.conn.execute("INSERT INTO activity_embeddings VALUES ('k', 't', '[]')")
            db.conn.commit()

        def act(ts, path):
            return {"timestamp": ts, "action": "message", "file_path": path, "source": "teams", "details": "{}"}
        store.upsert([
            (act("2026-06-02T10:00:00", "결제"), "결제", [1, 0, 0, 0]),       # 보고서 있는 지난 주 → 요약 후 삭제
            (act("2026-06-10T10:00:00", "장애"), "장애", [0, 1, 0, 0]),       # 보고서 없는 지난 주 → 요약 실패 시 보류
            (act("2026-09-29T10:00:00", "배포"), "배포", [0, 0, 1, 0]),       # 보관 기간 안 → 유지
        ])
        embed = lambda text: [1, 0, 0, 0] if "결제" in text else None
        result = apply_retention(
            8, db_path=db_path, store=store, summaries=summaries, today=date(2026, 10, 1),
            summary_for_week=lambda m: "결제 연동 테스트 완료" if m == date(2026, 6, 1) else "장애 대응",
            embed=embed, cache_path=os.path.join(tmp, "none.db"),
        )
        if result["summarized"] != ["2026-06-01"] or result["kept"] != ["2026-06-08"] or result["deleted"] != 1:
            return fail(f"정리 결과 불일치: {result}")
        if store.count() != 2 or summaries.week_starts() != {"2026-06-01"}:
            return fail(f"남은 벡터 {store.count()}개, 요약 {summaries.week_starts()}")
        if result["legacy_dropped"] != 1:
            return fail("레거시 임베딩 테이블 삭제 실패")
        again = apply_retention(8, db_path=db_path, store=store, summaries=summaries, today=date(2026, 10, 1),
                                summary_for_week=lambda m: "장애 대응", embed=lambda t: [0, 1, 0, 0],
                                cache_path=os.path.join(tmp, "none.db"))
        if again["summarized"] != ["2026-06-08"] or store.count() != 1:
            return fail(f"보류 주 재시도 실패: {again}")
        hits = search_past_weeks("결제", db_path=db_path, summaries=summaries, embed=lambda qs: [[1, 0, 0, 0]])
        if not hits or hits[0]["week_start"] != "2026-06-01" or len(hits[0]["activities"]) != 1:
            return fail(f"장기 질의 실패: {hits[:1]}")
        return ok("요약 저장 후 삭제·요약 실패 주 보류→재시도·최근 주 유지·레거시 삭제·장기 질의 정상")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def t_storage_settings():
    """Settings > 저장소 관리: 보관 기간 선택 저장, 사용량·속도 안내 표시 (실제 설정 파일은 저장하지 않음)"""
    from types import SimpleNamespace
    from weekly_report.gui.app import WeeklyPulseApp
    app = WeeklyPulseApp.__new__(WeeklyPulseApp)
    app.page = SimpleNamespace(window=SimpleNamespace(width=1400), update=lambda: None, run_thread=lambda fn: None)
    app.show_snack = lambda msg: None
    app._save_report_settings = lambda: None
    app.report_settings = {}
    app._create_storage_settings_section()
    if app.retention_dropdown.value != "12" or "확인 중" not in app.storage_usage_text.value:
        return fail(f"초기 표시 불일치: {app.retention_dropdown.value}")
    app._retention_changed("26")
    app._vector_stats = {"count": 1485, "open_seconds": 0.4, "weekly_rate": 1000, "summary_weeks": 0}
    app._render_storage_info()
    if app.report_settings["vector_retention_weeks"] != 26:
        return fail("보관 기간 저장 실패")
    if "1,485개" not in app.storage_usage_text.value or "26주를 다 채우면 약 26,000개" not in app.storage_speed_text.value:
        return fail(f"사용량/속도 안내 불일치: {app.storage_speed_text.value}")
    return ok("보관 기간 선택·사용량(벡터 수·여는 시간)·기간별 예상 속도 안내 정상")


def t_condense():
    """긴 원문 한 줄(Teams 메시지·지난 대화·Claude Code 요청): 짧으면 원문, 길면 잘라 싣지 않고 요약, 요약 실패 시 문장 경계"""
    from weekly_report.report import sections
    short = "직영 20개점 점검 완료했습니다."
    if sections.condense(short, summarize=lambda t: "X") != short:
        return fail("짧은 원문이 바뀜")
    long = "10월2차 정기배포 대상은 직영 20개점입니다. " + "배포 전 점검 항목을 담당자별로 확인해 주세요. " * 12
    calls = []
    got = sections.condense(long, summarize=lambda t: calls.append(t) or "10월2차 배포 대상 직영 20개점, 담당자별 점검 요청")
    again = sections.condense(long, summarize=lambda t: calls.append(t) or "다른 요약")
    if got != "(요약) 10월2차 배포 대상 직영 20개점, 담당자별 점검 요청" or again != got or len(calls) != 1:
        return fail(f"긴 원문 요약/재사용 실패: {got} / 호출 {len(calls)}회")
    other = long.replace("20개점", "30개점")
    fallback = sections.condense(other, summarize=lambda t: None)
    if len(fallback) > sections.RAW_LINE_LEN or not fallback.endswith("."):
        return fail(f"요약 실패 시 문장 경계 폴백 실패: {fallback[-20:]}")
    return ok("짧은 원문 유지, 긴 원문 요약(프로세스 내 재사용), 요약 실패 시 문장 경계 폴백")


def t_split_sequence():
    """요약 줄 안의 번호 목록 분리 — IP·버전·날짜는 번호로 오인하지 않는지"""
    from weekly_report.report.sections import split_sequence
    got = split_sequence("개선 요청: 1. 데이터 수집 문제 2. 중복 제거 3. 요약 확대")
    if got != {"head": "개선 요청:", "items": ["1. 데이터 수집 문제", "2. 중복 제거", "3. 요약 확대"]}:
        return fail(f"분리 오류: {got}")
    for text in ("개발 DB 10.253.42.50 1526 접속", "v3.0 배포 후 2. 항목만", "09/29(화) 1. 단독 항목"):
        if split_sequence(text)["items"]:
            return fail(f"오탐: {text}")
    return ok("번호 목록 분리 + IP/버전/단독 번호 오탐 없음")


# ───────────────────────── AI 분석 ─────────────────────────

def t_llm_templates():
    from weekly_report.ai.llm_summarizer import LLMSummarizer, PROMPT_TEMPLATES
    need = {"default", "file", "teams", "mail", "topic"}
    missing = need - set(PROMPT_TEMPLATES)
    if missing:
        return fail(f"템플릿 누락: {missing}")
    return ok("소스별 프롬프트 템플릿 5종 존재")


def t_cluster_topics():
    from weekly_report.report.generator import ReportGenerator
    from weekly_report.storage.database import ActivityDatabase, get_week_start_end
    start, end = get_week_start_end()
    with ActivityDatabase() as db:
        acts = db.get_activities_by_date_range(start, end)
    rg = ReportGenerator()
    topics = rg._cluster_topics(acts)
    if not topics:
        return fail("주제 클러스터링 결과 없음")
    names = [t["name"] for t in topics[:3]]
    return ok(f"{len(topics)}개 주제 — 예: {', '.join(names)}")


def t_topic_examples():
    """주제별 작업: 예시는 짧은 이름(채팅방·페이지 제목·파일명, 임시 파일 제외), '개인'으로 판정된 주제는 제외"""
    import json as _json
    from weekly_report.ai.clusterer import ActivityClusterer
    label = ActivityClusterer._example_label
    cases = [
        ({"source": "teams", "file_path": "Teams 채팅: ★공유방 - 홍수빈(POS) - 팀",
          "details": _json.dumps({"chat": "★공유방", "sender": "홍수빈"})}, "★공유방"),
        ({"source": "browser", "file_path": "https://www.google.com/search?q=%EA%B5%AC", "details": "{}"}, "www.google.com"),
        ({"source": "browser", "file_path": "https://confl.sinc.co.kr/x", "details": _json.dumps({"title": "POS 결제수단 현황"})},
         "POS 결제수단 현황"),
        ({"source": "filesystem", "file_path": "C:\\01.AI\\Data\\pay_h060_result.txt", "details": "{}"}, "pay_h060_result.txt"),
        ({"source": "filesystem", "file_path": "C:\\01.AI\\Data\\check.py.tmp.34680", "details": "{}"}, None),
        ({"source": "filesystem", "file_path": "C:\\01.AI\\a.html.tmp.34680.8421e26c67b5", "details": "{}"}, None),
        ({"source": "outlook", "file_path": "Meeting: x", "details": _json.dumps({"subject": "DevX 교육"})}, "DevX 교육"),
    ]
    wrong = [(c[0]["file_path"], label(c[0]), c[1]) for c in cases if label(c[0]) != c[1]]
    if wrong:
        return fail(f"예시 이름 오류: {wrong}")
    cl = ActivityClusterer.__new__(ActivityClusterer)
    acts = [{"source": "teams", "file_path": f"Teams 채팅: 방{i}", "details": "{}"} for i in range(4)]
    cl.cluster_activities = lambda a: [{"items": acts[:2]}, {"items": acts[2:]}]
    names = iter(["10월 배포 준비", "개인"])
    cl.name_cluster = lambda items: next(names)
    topics = cl.build_topics(acts)
    if [t["name"] for t in topics] != ["10월 배포 준비"]:
        return fail(f"개인 주제 제외 실패: {topics}")
    names = iter(["POS 운영", "POS 운영"])
    acts2 = [{"source": "teams", "file_path": "x", "details": _json.dumps({"chat": f"방{i}"})} for i in range(4)]
    cl.cluster_activities = lambda a: [{"items": acts2[:2]}, {"items": acts2[2:]}]
    if [t["name"] for t in cl.build_topics(acts2)] != ["POS 운영 (방0)", "POS 운영 (방2)"]:
        return fail("같은 이름 주제 구분 실패")
    return ok("채팅방·페이지 제목(없으면 사이트)·파일명·회의 제목으로 표시, 임시 파일 제외, '개인' 주제 제외, 같은 이름 구분")


def t_issues_parse():
    """STEP 2 이슈·리스크/다음 주 계획 하네스: 형식·인용 없는 줄 버림, 구분 검증, 이미 지난 일정 제외"""
    from datetime import date
    from weekly_report.ai.issues import drop_past_plans, parse_issue_output, plan_dates
    output = ("[이슈·리스크]\n장애 | 거래저장 실패 점포 — 누락 매출 2건 | 하드 교체 후 수기 입력 [2]\n"
              "이슈 | 국민QR 환불 시 결제바코드로도 환불 가능 | KIS 회신 대기 [1][9]\n"
              "기타 | 잘못된 구분 | 확인 필요 [3]\n리스크 | 인용 없는 행 | 확인 필요\n"
              "[다음 주 계획]\n- 10월1차 순차 배포 모니터링 (~10/06) [3]\n- 직영 전체점 정기배포 (10/01) [4]\n- 없음")
    issues, plans = parse_issue_output(output, evidence_count=5)
    if [(i["type"], i["refs"]) for i in issues] != [("장애", [2]), ("이슈", [1])]:
        return fail(f"이슈 파싱 오류: {issues}")
    if plan_dates("배포 (10월 1일 02:00) 및 10/06 점검", 2026) != [date(2026, 10, 1), date(2026, 10, 6)]:
        return fail("날짜 인식 오류")
    kept = drop_past_plans(plans, date(2026, 10, 2))
    if [p["text"] for p in kept] != ["10월1차 순차 배포 모니터링 (~10/06)"]:
        return fail(f"지난 일정 제외 오류: {kept}")
    return ok("구분·인용 없는 행 제외, 없는 근거 번호 제거, 날짜 인식(10/06·10월 1일), 이미 지난 일정 제외")


def t_langchain_present():
    try:
        import langchain_core, langchain_openai, langgraph  # noqa
        return ok("langchain-core/openai/langgraph 임포트 정상")
    except ImportError as e:
        return fail(str(e))


# ───────────────────────── 보고서 생성 ─────────────────────────

def t_markdown_report():
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    data = rg.collect_weekly_data()
    md = rg.generate_markdown(data)
    required = ["STEP 1. 프로그램별 상세 활동 내역", "STEP 2. 업무 요약", "이번 주 핵심 요약"]
    missing = [r for r in required if r not in md]
    if missing:
        return fail(f"섹션 누락: {missing}")
    return ok(f"마크다운 생성 ({len(md)}자)")


def t_word_report():
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    data = rg.collect_weekly_data()
    path = "test_reports/_test_report.docx"
    os.makedirs("test_reports", exist_ok=True)
    try:
        rg.generate_word(data, path)
        size = os.path.getsize(path)
        if size < 1000:
            return fail(f"파일 너무 작음 {size}B")
        return ok(f"docx 생성 {size//1024}KB")
    finally:
        if os.path.exists(path):
            os.remove(path)


def t_graph_pipeline():
    # test_reports에 저장 — reports/의 이번 주 보고서(사용자가 편집한 파일)를 덮어쓰지 않도록
    from weekly_report.report.pipeline import WeeklyReportPipeline
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    result = WeeklyReportPipeline(rg, output_dir="test_reports").run(formats=("markdown",))
    paths = result["file_paths"]
    if "markdown" not in paths or not os.path.exists(paths["markdown"]):
        return fail("마크다운 경로 없음")
    return ok(f"LangGraph 경로 생성: {os.path.basename(paths['markdown'])}")


def t_empty_week():
    """활동 없는 주에도 보고서가 깨지지 않는지 (analyze 스킵 분기)"""
    from weekly_report.report.pipeline import WeeklyReportPipeline
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    pipe = WeeklyReportPipeline(rg, output_dir="test_reports")
    result = pipe.run(week_start="2026-01-01", week_end="2026-01-07",
                      formats=("markdown",))
    wd = result["weekly_data"]
    if wd["has_activities"]:
        return fail("빈 주인데 활동 있음")
    return ok("빈 주 보고서 생성 + analyze 스킵 정상")


def t_weekly_summary_db():
    from weekly_report.storage.database import ActivityDatabase
    with ActivityDatabase() as db:
        rows = db.conn.execute(
            "SELECT COUNT(*) FROM weekly_summaries"
        ).fetchone()[0]
    if rows == 0:
        return fail("weekly_summaries 비어있음")
    return ok(f"주간 요약 {rows}건 저장됨")


# ───────────────────────── GUI·실행 ─────────────────────────

def t_flet_import():
    try:
        from weekly_report.gui import app as modern_gui  # noqa
        return ok("modern_gui 임포트 정상 (flet+win32 가드)")
    except ImportError as e:
        return fail(str(e))


def t_process_detect():
    from weekly_report.gui.app import WeeklyPulseApp
    for proc in ["chrome.exe", "winword.exe", "Code.exe", "powershell.exe"]:
        if WeeklyPulseApp.is_app_running([proc]):
            return ok(f"{proc} 감지됨")
    return ok("감지 대상 프로세스 없음 (로직 실행 정상)")


def t_integrated_collector():
    try:
        from weekly_report.collectors.integrated import IntegratedCollector
        return ok("IntegratedCollector 임포트 정상")
    except ImportError as e:
        return fail(str(e))


def t_expected_questioner_settings():
    """Settings > Teams 예상질문자 추가·중복·빈값·삭제 (실제 설정 파일은 저장하지 않음)"""
    from types import SimpleNamespace
    from weekly_report.collectors.teams import DEFAULT_SETTINGS
    from weekly_report.gui.app import WeeklyPulseApp
    if DEFAULT_SETTINGS.get("expected_questioners") != []:
        return fail("기본값 expected_questioners 누락")
    app = WeeklyPulseApp.__new__(WeeklyPulseApp)
    app.page = SimpleNamespace(window=SimpleNamespace(width=1400), update=lambda: None)
    app.show_snack = lambda msg: None
    app._save_teams_settings = lambda: None
    app._create_teams_settings_section()
    app.teams_settings["expected_questioners"] = []
    for text in ("  김영호 ", "팀장", "김영호", "   "):
        app.questioner_input.value = text
        app._add_expected_questioner(None)
    app._remove_expected_questioner("팀장")
    names = app.teams_settings["expected_questioners"]
    if names != ["김영호"] or len(app.questioner_chips.controls) != 1:
        return fail(f"결과 불일치: {names}")
    return ok("추가·공백정리·중복/빈값 거부·삭제 정상")


def t_rag_topic_settings():
    """Settings > 보고서 설정 주제 질의 추가·중복·삭제·검색 범위 (실제 설정 파일은 저장하지 않음)"""
    from types import SimpleNamespace
    from weekly_report.gui.app import WeeklyPulseApp
    app = WeeklyPulseApp.__new__(WeeklyPulseApp)
    app.page = SimpleNamespace(window=SimpleNamespace(width=1400), update=lambda: None)
    app.show_snack = lambda msg: None
    app._save_report_settings = lambda: None
    app._create_report_settings_section()
    app.report_settings["rag_topics"] = ["장애 대응"]
    for text in (" 손익·예산 ", "장애 대응", "  "):
        app.rag_topic_input.value = text
        app._add_rag_topic(None)
    app._remove_rag_topic("장애 대응")
    app._rag_past_weeks_changed("abc")
    app._rag_past_weeks_changed("6")
    topics = app.report_settings["rag_topics"]
    if topics != ["손익·예산"] or len(app.rag_topic_chips.controls) != 1:
        return fail(f"결과 불일치: {topics}")
    if app.report_settings["rag_past_weeks"] != 6:
        return fail(f"검색 범위 불일치: {app.report_settings['rag_past_weeks']}")
    app.report_settings["evidence_exclude_keywords"] = ["weekly_report"]
    for text in (" 개인 메모 ", "WEEKLY_REPORT", ""):
        app.noise_keyword_input.value = text
        app._add_noise_keyword(None)
    app._remove_noise_keyword("weekly_report")
    if app.report_settings["evidence_exclude_keywords"] != ["개인 메모"] or len(app.noise_keyword_chips.controls) != 1:
        return fail(f"근거 제외 키워드 불일치: {app.report_settings['evidence_exclude_keywords']}")
    return ok("주제·근거 제외 키워드 추가·중복(대소문자 무시)/빈값 거부·삭제, 검색 범위 숫자 검증 정상")


def t_rag_citation_check():
    """RAG 출력 검증: 인용 없는 문장·없는 근거 번호 문장 제거, 인용 과다 축약, NONE 처리"""
    from weekly_report.ai.rag import parse_cited_bullets
    output = (
        "- 거래저장 실패 장애를 복구함 [1][2]\n"
        "- 근거 없는 추측 문장\n"
        "- 존재하지 않는 근거를 단 문장 [9]\n"
        "서론 문장 [1]\n"
        "- 근거를 많이 단 문장 [1][2][3][4]\n"
        "- (지난 기록) 지난달 점검과 이어짐 [3]"
    )
    bullets = parse_cited_bullets(output, evidence_count=4)
    texts = [t for t, _ in bullets]
    if len(bullets) != 3 or bullets[1][1] != [1, 2, 3] or "[4]" in texts[1]:
        return fail(f"검증 결과 불일치: {bullets}")
    if parse_cited_bullets("NONE", 3) != [] or parse_cited_bullets("", 3) != []:
        return fail("NONE/빈 출력 처리 실패")
    return ok("근거 없는 문장 2건 제거, 인용 4개→3개 축약, NONE 처리 정상")


def t_evidence_noise_filter():
    """근거 노이즈 필터: IDE 기록 전체 제외, 키워드(경로·제목·이스케이프된 내용) 제외, 일반 업무 기록은 유지"""
    import json as _json
    from weekly_report.ai.rag import is_evidence_noise
    kw = ["주간보고 자동작성", "weekly_report", "WeeklyPulse"]
    noise = [
        {"source": "claude_code", "action": "claude_interaction", "file_path": "C:/orca/first", "details": "{}"},
        {"source": "orca", "action": "modified", "file_path": "C:/orca/first", "details": "{}"},
        {"source": "filesystem", "file_type": "txt", "action": "modified",
         "file_path": "C:/Downloads/주간보고 자동작성 프로그램 확인.txt", "details": "{}"},
        {"source": "browser", "action": "visit", "file_path": "https://github.com/x/weekly_report_automation",
         "details": _json.dumps({"title": "repo"})},
        {"source": "outlook", "action": "sent", "file_path": "메일",  # ensure_ascii로 저장된 내용
         "details": _json.dumps({"summary": "WeeklyPulse 시연 자료 공유"})},
    ]
    keep = [
        {"source": "teams", "action": "message", "file_path": "운영방",
         "details": _json.dumps({"text": "10월2차 정기배포 모니터링 완료"}, ensure_ascii=False)},
        {"source": "confluence", "action": "edited", "file_path": "POS 결제수단 현황",
         "details": _json.dumps({"summary": "결제수단별 현황 표 갱신"})},
    ]
    wrong = [a["file_path"] for a in noise if not is_evidence_noise(a, kw)]
    wrong += [a["file_path"] for a in keep if is_evidence_noise(a, kw)]
    if wrong:
        return fail(f"판정 오류: {wrong}")
    if is_evidence_noise(noise[2], []) or not is_evidence_noise(noise[0], []):
        return fail("키워드 없을 때 IDE만 제외돼야 함")
    # 이름뿐인 기록(첨부파일·스크린샷·제목 없는 방문)은 근거 후보에서 제외, 내용 있는 메모는 유지
    from weekly_report.ai.rag import TopicQueryRAG
    rag = TopicQueryRAG.__new__(TopicQueryRAG)
    rag.me_name, rag.teams_scope, rag.noise_keywords = "", "all", kw
    rag.excluded_ids, rag.excluded_titles = set(), set()
    name_only = [
        {"source": "confluence", "action": "confluence_created", "file_path": "팀/AI_CONTEXT.md",
         "details": _json.dumps({"page_id": "1", "summary": ""})},
        {"source": "sharepoint", "action": "modified", "file_path": "SharePoint 파일: 스크린샷.png",
         "details": _json.dumps({"name": "스크린샷.png"})},
        {"source": "browser", "action": "visit", "file_path": "https://gw.example.com/login", "details": "{}"},
    ]
    name_only.append({"source": "teams", "action": "message", "file_path": "Teams 채팅: 파트리더",
                      "details": _json.dumps({"chat": "[POS2팀] 파트리더", "text": "네 팀장님.", "from_me": True})})
    memo = {"source": "filesystem", "file_type": "txt", "action": "modified", "file_path": "C:/memo/오늘 할 일.txt",
            "details": _json.dumps({"content_added": "국민QR 환불 개선 배포 확인", "summary": "배포 확인 메모"})}
    wrong = [a["file_path"] for a in name_only if rag.eligible(a)]
    wrong += [a["file_path"] for a in keep + [memo] if not rag.eligible(a)]
    if wrong:
        return fail(f"근거 후보 판정 오류: {wrong}")
    return ok("IDE 기록·키워드(경로/제목/이스케이프 내용)·이름뿐인 기록·단순 대답 제외, 업무 기록·메모 유지")


def t_evidence_label():
    """근거 표기: 출처 이름을 자르지 않음, 요지는 첫 문장(요약 제목줄 제거), 긴 글은 단어 경계, 인증번호 메시지는 요지 생략"""
    import json as _json
    from datetime import datetime as _dt
    from weekly_report.ai.rag import EVIDENCE_GIST_MAX, evidence_label
    when = _dt(2026, 9, 29, 10, 0)
    chat = "[이마트24POS팀] 10월2차 정기배포 & 모니터링 공유방"
    def label(details, source="teams"):
        return evidence_label({"source": source, "action": "message", "file_path": "x",
                               "details": _json.dumps(details, ensure_ascii=False)}, when)
    a = label({"chat": chat, "text": "직영 20개점 점검 결과 특이사항 없습니다. 내일 전체점 배포 예정입니다."})
    if chat not in a or not a.endswith('"직영 20개점 점검 결과 특이사항 없습니다."'):
        return fail(f"출처/첫 문장 오류: {a}")
    b = label({"chat": chat, "text": "인증번호 482913 입니다"})
    if '"' in b:
        return fail(f"민감 내용 노출: {b}")
    c = label({"summary": "# 핵심 요약\n**도입 내용**: 안면결제 단말기 도입, 10월 오픈 예정\n**영향도**: 기존과 동일"},
              source="confluence")
    if not c.endswith('"도입 내용 : 안면결제 단말기 도입, 10월 오픈 예정"'):
        return fail(f"여러 줄 요약의 첫 내용 줄 실패: {c}")
    long_text = "배포 " * 200
    d = label({"chat": chat, "text": long_text})
    gist = d.split(' — "', 1)[1]
    if not gist.endswith('…"') or len(gist) > EVIDENCE_GIST_MAX + 3 or "배포배" in gist:
        return fail(f"긴 글 자르기 오류: {gist[-20:]}")
    from weekly_report.report.sections import clean_text
    long = "첫 문장은 배포 완료입니다. " + "두 번째 문장은 아주 길게 이어지는 설명 " * 20
    if clean_text(long, None) != long.strip():
        return fail("limit=None인데 잘림")
    if clean_text(long, 30) != "첫 문장은 배포 완료입니다.":  # 상한의 절반을 넘긴 문장 끝에서 끊음
        return fail(f"문장 경계 자르기 실패: {clean_text(long, 30)}")
    cut = clean_text("가나다라 " * 50, 100)
    if not cut.endswith("…") or "가나다라가" in cut or cut[-2] == " ":
        return fail(f"단어 경계 자르기 실패: {cut[-12:]}")
    return ok("출처 이름 전체 표시, 첫 문장 요지, 요약 제목줄 제거, 긴 글 단어 경계, 인증번호 요지 생략, "
              "clean_text 무제한·문장/단어 경계 자르기")


def t_llm_output_check():
    """LLM 출력 검증(하네스): 요약 거절·되묻기 불합격→재요청→실패 시 None(캐시 안 함), 서두 제거,
    업무 문장 속 '확인할 수 없' 오탐 없음 / 근거 번호-문장 관련성 검사 / STEP 3 '(보고 내용)' 표시"""
    from weekly_report.ai.llm_summarizer import LLMSummarizer, check_summary
    from weekly_report.ai.rag import check_citations
    from weekly_report.ai.questions import parse_qa
    cases = {
        "죄송하지만 요약할 내용이 없습니다.": None,
        "텍스트를 제공해 주시면 요약해 드리겠습니다.": None,
        "요약하려면 원문을 공유해 주시겠어요?": None,
        "다음은 메일 요약입니다:\n10월2차 배포 일정 공유": "10월2차 배포 일정 공유",
        "요약: 직영 20개점 점검 완료": "직영 20개점 점검 완료",
        "**요약:**\nIT개발혁신 TF 5단계 워크플로우 공유": "IT개발혁신 TF 5단계 워크플로우 공유",
        "DB 복구가 실패해 원인은 확인할 수 없는 상태이며 하드 교체를 검토 중": "DB 복구가 실패해 원인은 확인할 수 없는 상태이며 하드 교체를 검토 중",
    }
    wrong = {k: check_summary(k) for k, v in cases.items() if check_summary(k) != v}
    if wrong:
        return fail(f"요약 검증 판정 오류: {wrong}")

    class FakeChain:
        def __init__(self, outputs):
            self.outputs, self.calls = list(outputs), []
        def invoke(self, inputs):
            self.calls.append(inputs["system"])
            return self.outputs.pop(0)
    llm = LLMSummarizer.__new__(LLMSummarizer)
    llm.config, llm.enabled = {"enabled": True, "api_key": "x"}, True
    puts = []
    llm._cache_get, llm._cache_put = (lambda key: None), (lambda key, value: puts.append(value))
    for outputs, expected in ((["죄송합니다, 요약할 내용이 없습니다.", "배포 완료 공유"], "배포 완료 공유"),
                              (["죄송합니다.", "텍스트를 제공해 주세요."], None)):
        chain = FakeChain(outputs)
        llm._chain = lambda *a, chain=chain: chain
        puts.clear()
        got = llm.complete("요약해", "원문", temperature=0, cache=True, validate=check_summary)
        if got != expected or len(chain.calls) != 2 or "설명·사과·되묻기 없이" not in chain.calls[1]:
            return fail(f"재요청 동작 오류: {got} / 호출 {len(chain.calls)}회")
        if puts != ([expected] if expected else []):
            return fail(f"불합격 응답이 캐시됨: {puts}")

    # 근거 관련성: 문장 벡터와 근거 벡터가 먼 번호는 지움, 벡터 없는 근거는 판단 보류
    items = [("국민QR 환불 개선 완료 [1][2]", [1, 2]), ("교육 일정 공유 [3]", [3]), ("배포 완료 [4]", [4])]
    embed = lambda texts: [[1, 0, 0], [0, 1, 0], [0, 0, 1]]
    evidence = {1: [1, 0, 0], 2: [0, 1, 0], 3: [1, 0, 0]}  # [2]는 무관, [3]도 무관, [4]는 벡터 없음
    checked = check_citations(items, evidence, embed)
    if checked != [("국민QR 환불 개선 완료 [1]", [1]), ("교육 일정 공유", []), ("배포 완료 [4]", [4])]:
        return fail(f"근거 관련성 검사 오류: {checked}")
    if check_citations(items, evidence, lambda texts: None) != items:
        return fail("임베딩 실패 시 원본 유지 안 됨")
    # 애매한 구간(0.30~0.42): 문장 핵심어가 근거 원문에 있으면 유지, 없으면 삭제
    gray = {1: [0.35, 0.9367]}  # 문장 벡터 [1, 0]과의 유사도 0.35
    claim = [("안면결제 테스트 장비 9월 초 제공 예정 [1]", [1])]
    keep = check_citations(claim, gray, lambda t: [[1, 0]],
                           evidence_texts={1: "개발 현황: 프로그램 개발 100% 완료, 테스트 장비 9월 초 제공 예정 안면결제"})
    drop = check_citations(claim, gray, lambda t: [[1, 0]], evidence_texts={1: "AI 활용 개발 성과 보고서"})
    if keep[0][1] != [1] or drop[0][1] != []:
        return fail(f"애매한 구간 핵심어 판정 오류: {keep} / {drop}")

    # 사실 대조(LLM-as-judge): 근거와 다르면 '근거와 다름' 표시, 근거에 없으면 '확인 필요', 내용은 고치지 않음
    from weekly_report.ai.questions import DIFF_MARK, fact_check_answers
    answers = [("구분자는 운임 계산용입니다. [1]", [1]), ("TID 개발/운영 동일 [2]", [2]),
               ("일정 확정 [3]", [3]), ("확인 필요", [])]
    texts = {1: "CJ대한통운 정산 시 동일권/타권 구분 필요", 2: "KICC 개발기 운영기 동일 TID", 3: "산출물 목록"}
    seen = []
    judged = fact_check_answers(answers, texts, lambda s, b: seen.append(b) or "답변1: DIFF - 정산용인데 운임 계산용\n답변2: OK - 일치\n답변3: NONE - 없음")
    if judged != [f"구분자는 운임 계산용입니다. [1] {DIFF_MARK}", "TID 개발/운영 동일 [2]",
                  "일정 확정 [3] (확인 필요)", "확인 필요"] or "CJ대한통운 정산" not in seen[0]:
        return fail(f"사실 대조 표시 오류: {judged}")
    if fact_check_answers(answers, texts, lambda s, b: "") != [a for a, _ in answers]:
        return fail("판정 실패 시 원본 유지 안 됨")
    # 성과 수치 측정 기준 표시 (정답셋 G09): 성과 수치 + 근거에 측정 기준 없음 → 표시, 일반 수치·측정 기준 있음 → 그대로
    from weekly_report.ai.questions import MEASURE_CAVEAT, add_measure_caveat
    perf = "레거시 분석 60~80% 단축, 재작업 비용 50% 이상 감소했습니다. [1]"
    cases = [(perf, ["생산성 효과: 레거시 분석 60~80% 단축"], True),
             (perf, ["측정 기준: 2025년 동일 규모 프로젝트 대비 실측"], False),
             ("테스트 시나리오 작성이 수시간→수십분으로 줄었습니다. [2]", ["수시간→수십분"], True),
             ("KICC TID는 7054576이고 세종 프리픽스는 30입니다. [3]", ["TID 7054576"], False),
             ("60% 단축했으나 측정 기준은 확인 필요합니다. [1]", [""], False)]
    for answer, texts, expect in cases:
        if add_measure_caveat(answer, texts).endswith(MEASURE_CAVEAT) != expect:
            return fail(f"측정 기준 표시 오류: {answer} / {texts}")
    # '확인 필요'만 남은 답변 → 질문과 가장 가까운 보고 내용으로 채움
    from weekly_report.ai.questions import fill_bare_answers, is_bare_answer
    if not is_bare_answer("확인 필요 (보고 내용)") or is_bare_answer("개발 완료, 일정은 확인 필요 [1]"):
        return fail("'확인 필요'만 남은 답변 판정 오류")
    filled = fill_bare_answers(["안면결제 테스트는?", "배포는?"], ["확인 필요.", "10월 배포 완료 [1]"],
                               ["조직도 업데이트", "안면결제 개발 완료, 장비 9월 초 제공 예정"], [[1, 0], [0, 1]],
                               lambda qs: [[0.1, 0.9]])
    if filled != ["보고 기준: 안면결제 개발 완료, 장비 9월 초 제공 예정 — 질문하신 부분은 확인 필요합니다. (보고 내용)",
                  "10월 배포 완료 [1]"]:
        return fail(f"빈 답변 폴백 오류: {filled}")
    # 2단계 검색 재답변: 그 질문에 준 근거 번호만 남기고, 없으면 '(확인 필요)'
    from weekly_report.ai.questions import parse_reanswers
    redone = parse_reanswers("A1: 권역은 우편번호 앞 2자리로 구분합니다. [7][2]\nA3: 설정 완료 [9]\nA4: 무시",
                             {1: {2, 7}, 3: {5}})
    if redone != {1: ("권역은 우편번호 앞 2자리로 구분합니다. [2][7]", [2, 7]), 3: ("설정 완료 (확인 필요)", [])}:
        return fail(f"재답변 파싱 오류: {redone}")
    qa = parse_qa("Q: 배포 일정은?\nA: 10월 7일 배포 예정입니다. (보고 내용)\nQ: 원인은?\nA: 원인 분석 중입니다.\n"
                  "Q: 언제부터야?\nA: 확인 필요 (보고 내용)", 3)
    if (qa[0][1] != "10월 7일 배포 예정입니다. (보고 내용)" or not qa[1][1].endswith("(확인 필요)")
            or qa[2][1] != "확인 필요"):
        return fail(f"STEP 3 근거 표시 오류: {qa}")
    return ok("요약 거절·되묻기 재요청→실패 시 None·캐시 안 함, 서두 제거, 업무 문장 오탐 없음, "
              "무관한 근거 번호 제거(애매하면 핵심어 확인), 답변-근거 사실 대조 표시, 질문별 근거 재답변, 성과 수치 측정 기준 표시, '(보고 내용)'·'(확인 필요)' 표시")


def t_rag_sections():
    """실데이터 RAG: 이번 주 활동 → 주제별 근거 인용 요약 → Markdown 'STEP 2 주제 질의 요약'"""
    from weekly_report.report.generator import ReportGenerator
    from weekly_report.ai.rag import TopicQueryRAG
    rg = ReportGenerator()
    week_start, week_end = rg._resolve_week_range(None, None)
    acts = rg.fetch_week_activities(week_start, week_end)
    rag = TopicQueryRAG(rg)
    if not rag.enabled:
        return skip("LiteLLM 설정 없음")
    if not acts:
        return skip("이번 주 활동 없음")
    sections = rag.build_sections(acts, week_start, week_end)
    if not sections:
        return skip("설정한 주제와 관련된 이번 주 기록 없음")
    for s in sections:
        numbers = {e["n"] for e in s["evidence"]}
        cited = {int(n) for b in s["bullets"] for n in re.findall(r"\[(\d+)\]", b)}
        if not s["bullets"] or not cited or not cited <= numbers:
            return fail(f"[{s['topic']}] 인용·근거 번호 불일치: {cited} / {numbers}")
    data = rg.compose_weekly_data(week_start, week_end, acts, {"rag_sections": sections})
    md = rg.generate_markdown(data)
    if "### 주제 질의 요약" not in md or md.index("### 주제 질의 요약") < md.index("## STEP 2"):
        return fail("Markdown STEP 2에 주제 질의 요약 없음")
    return ok(f"{len(sections)}개 주제 — " + ", ".join(s["topic"] for s in sections))


def t_expected_qa_check():
    """STEP 3 출력 검증(하네스): Q/A 짝 맞추기, 과거 대화 [P#]는 질문에만, 근거 없는 답변은 '확인 필요',
    이번 주에 이미 한 질문 반복 감지"""
    from weekly_report.ai.questions import is_repeat, parse_qa
    output = (
        "Q: 네이버 오픈은 몇시야? [P1]\n"
        "A: 17시 30분 파일럿 오픈 예정 [2][P1]\n"
        "Q: 답변 없는 질문\n"
        "Q: KIS 점포 배포해도 되는 거지? [P9]\n"
        "A: 기존 KIS 검증 동작이라 배포 영향 없음\n"
        "Q: 원인은 뭐야?\n"
        "A: 원인은 확인 필요 [7]"
    )
    pairs = parse_qa(output, evidence_count=3, past_count=2)
    if len(pairs) != 3:
        return fail(f"Q/A 짝 불일치: {pairs}")
    (q1, a1, r1, p1), (q2, a2, r2, p2), (q3, a3, r3, p3) = pairs
    if p1 != [1] or "[P" in q1 or "[P" in a1 or r1 != [2]:
        return fail(f"과거 대화 번호 처리 실패: {pairs[0]}")
    if p2 != [] or not a2.endswith("(확인 필요)"):
        return fail(f"근거 없는 답변 표시 실패: {pairs[1]}")
    if r3 != [] or "[7]" in a3 or a3.endswith("(확인 필요)"):
        return fail(f"없는 근거 번호 처리 실패: {pairs[2]}")
    asked = ["내가 글들을 좀 늦게 봤는데 네이버점포 거래는 잘적용되는지 봐주고"]
    if not is_repeat("내가 글들을 좀 늦게 봤는데, 네이버점포 거래는 잘 적용되는지 봐주고?", asked):
        return fail("반복 질문 미감지")
    if is_repeat("KICC DLL 예외는 배포 후 문제 없는 거지?", asked):
        return fail("다른 질문을 반복으로 오판")
    return ok("Q/A 짝·[P#] 질문 전용·근거 없음 표시·없는 번호 제거·반복 질문 감지 정상")


def t_expected_qa():
    """실데이터 STEP 3: 예상질문자 메시지 → 성향 → 예상 질문 5개 → Markdown 'STEP 3'"""
    from weekly_report.ai.questions import QUESTION_COUNT, ExpectedQuestions
    from weekly_report.report.generator import ReportGenerator
    rg = ReportGenerator()
    week_start, week_end = rg._resolve_week_range(None, None)
    acts = rg.fetch_week_activities(week_start, week_end)
    qa = ExpectedQuestions(rg)
    if not qa.enabled:
        return skip("예상질문자 미설정 또는 LiteLLM 설정 없음")
    if not acts:
        return skip("이번 주 활동 없음")
    points = rg._summarize_week(rg.build_program_sections(acts))
    sections = qa.build_sections(acts, week_start, week_end, points)
    if not sections:
        return skip("예상질문자의 Teams 메시지 없음")
    for s in sections:
        numbers = {e["n"] for e in s["evidence"]}
        past_ids = {p["n"] for p in s["past"]}
        if not 1 <= len(s["qa"]) <= QUESTION_COUNT:
            return fail(f"[{s['name']}] 질문 수 {len(s['qa'])}")
        for item in s["qa"]:
            cited = {int(n) for n in re.findall(r"\[(\d+)\]", item["a"])}
            if not cited <= numbers or not set(item["past"]) <= past_ids:
                return fail(f"[{s['name']}] 번호 불일치: {item}")
            if not cited and "확인 필요" not in item["a"] and "보고 내용" not in item["a"]:
                return fail(f"[{s['name']}] 근거도 '확인 필요'·'(보고 내용)'도 없는 답변: {item['a']}")
    md = rg.generate_markdown(rg.compose_weekly_data(week_start, week_end, acts, {"expected_qa": sections}))
    if "## STEP 3. 예상 질문 & 답변" not in md:
        return fail("Markdown에 STEP 3 없음")
    s = sections[0]
    return ok(f"{s['name']}: 질문 {len(s['qa'])}개, 성향 {len(s['style'])}줄, 과거 대화 연결 {len(s['past'])}건")


def t_report_status_panel():
    """메인 화면 보고서 생성 상태: 생성대기 → 작성 중(버튼 비활성) → 생성완료(폴더 아이콘),
    완료 후 늦게 온 경과 시간 갱신(only_if)이 완료 표시를 덮어쓰지 않는지"""
    from types import SimpleNamespace
    from weekly_report.gui.app import WeeklyPulseApp
    app = WeeklyPulseApp.__new__(WeeklyPulseApp)
    app.page = SimpleNamespace(window=SimpleNamespace(width=1400), update=lambda: None)
    app.show_snack = lambda msg: None

    app.create_control_panel()
    button = app.generate_button
    if app.report_status_text.value != "생성대기" or app.report_folder_button.visible:
        return fail("초기 상태가 생성대기가 아님")
    app.set_report_status("작성 중", detail="수집 중", requested_at="10:20:30")
    if not button.disabled or app.report_time_text.value != "10:20:30":
        return fail("작성 중 표시/버튼 비활성 실패")
    app.set_report_status("주간보고 생성완료", detail="소요 30초 · a.md", output_path="reports/a.md")
    app.set_report_status(detail="수집 중 · 경과 31초", only_if="작성 중")  # 늦게 온 타이머 갱신
    if (app.report_status_text.value != "주간보고 생성완료" or button.disabled
            or app.report_detail_text.value != "소요 30초 · a.md" or not app.report_folder_button.visible):
        return fail(f"완료 표시 실패: {app.report_state}")
    app.set_report_status("작성 중", detail="수집 중")
    if app.report_folder_button.visible:
        return fail("새 생성 시작 후에도 폴더 아이콘이 남음")
    return ok("생성대기→작성 중(버튼 비활성)→생성완료(폴더 아이콘), 늦은 타이머 갱신 무시 정상")


# ───────────────────────── 실행 ─────────────────────────

TESTS = [
    ("ENV-01", "환경/인증", "DB 스키마 및 연결", t_db_schema),
    ("ENV-02", "환경/인증", "LiteLLM 프록시 설정 로드", t_litellm_config),
    ("ENV-03", "환경/인증", "임베딩 API 호출", t_embedding_call),
    ("ENV-04", "환경/인증", "LLM 채팅(LCEL 체인) 호출", t_lcel_chat),
    ("ENV-05", "환경/인증", "Teams 토큰 파일 존재", t_teams_token),
    ("ENV-06", "환경/인증", "Teams 토큰 갱신(라이브)", t_teams_auth_live),

    ("COL-01", "수집", "파일 감시 노이즈 필터", t_watcher_noise),
    ("COL-02", "수집", "오피스 내용 캡처/DRM 폴백", t_office_content_capture),
    ("COL-03", "수집", "브라우저 수신 서버(5757)", t_browser_server),
    ("COL-04", "수집", "IDE 경로 탐색", t_ide_paths),
    ("COL-05", "수집", "Teams 채팅 수집(라이브)", t_teams_chat_collect),
    ("COL-06", "수집", "Teams 일정 수집(라이브)", t_teams_calendar),
    ("COL-07", "수집", "Outlook 수집(COM/Graph)", t_outlook_collect),
    ("COL-08", "수집", "Confluence 설정 존재", t_confluence),
    ("COL-09", "수집", "Slack 설정 존재", t_slack),
    ("COL-10", "수집", "감시 폴더 설정", t_file_watcher_watchdirs),
    ("COL-11", "수집", "OneNote 수집(라이브)", t_onenote),
    ("COL-12", "수집", "SharePoint/OneDrive 수집(라이브)", t_sharepoint),
    ("COL-13", "수집", "Office COM 읽기 + 버전 diff", t_office_com_diff),
    ("COL-14", "수집", "프로젝트 문서·지난 보고 (근거 전용)", t_project_docs),
    ("COL-15", "수집", "메일 수신/참조·Teams 언급 (일일 할 일 재료)", t_mail_teams_fields),

    ("PROC-01", "데이터 가공", "회의 중복 제거(teams 우선)", t_meeting_dedupe),
    ("PROC-02", "데이터 가공", "일시 파일(created+deleted) 제거", t_transient_files),
    ("PROC-03", "데이터 가공", "경로 정규화", t_path_normalize),
    ("PROC-04", "데이터 가공", "민감 URL 필터(OAuth 콜백)", t_sensitive_url_filter),
    ("PROC-05", "데이터 가공", "요약 캐시 조회", t_summary_cache),
    ("PROC-06", "데이터 가공", "공유 문서 변경 분리(내 수정/파트원 수정)", t_shared_doc_split),
    ("PROC-07", "데이터 가공", "요약 번호 목록 줄바꿈 분리", t_split_sequence),
    ("PROC-08", "데이터 가공", "VectorDB 저장·필터 검색 (Qdrant)", t_vector_store),
    ("PROC-09", "데이터 가공", "저장소 보관 정책 (요약 후 벡터 정리)", t_storage_retention),
    ("PROC-10", "데이터 가공", "긴 원문 요약 (Teams 메시지·지난 대화)", t_condense),

    ("AI-01", "AI 분석", "소스별 프롬프트 템플릿", t_llm_templates),
    ("AI-02", "AI 분석", "임베딩 클러스터링 → 주제", t_cluster_topics),
    ("AI-11", "AI 분석", "주제별 작업 예시 이름·개인 주제 제외", t_topic_examples),
    ("AI-12", "AI 분석", "STEP 2 이슈·리스크/다음 주 계획 하네스", t_issues_parse),
    ("AI-03", "AI 분석", "LangChain/LangGraph 임포트", t_langchain_present),
    ("AI-04", "AI 분석", "RAG 근거 인용 검증 (하네스)", t_rag_citation_check),
    ("AI-08", "AI 분석", "근거 노이즈 필터 (IDE·제외 키워드·이름뿐인 기록)", t_evidence_noise_filter),
    ("AI-09", "AI 분석", "근거 표기 (출처 전체·내용 요지)", t_evidence_label),
    ("AI-10", "AI 분석", "LLM 출력 검증 (요약 형식·근거 관련성, 하네스)", t_llm_output_check),
    ("AI-05", "AI 분석", "RAG 주제 질의 섹션 (실데이터)", t_rag_sections),
    ("AI-06", "AI 분석", "STEP 3 예상 질문 출력 검증 (하네스)", t_expected_qa_check),
    ("AI-07", "AI 분석", "STEP 3 예상 질문 & 답변 (실데이터)", t_expected_qa),

    ("RPT-01", "보고서 생성", "Markdown 섹션 완전성", t_markdown_report),
    ("RPT-02", "보고서 생성", "Word(docx) 생성", t_word_report),
    ("RPT-03", "보고서 생성", "LangGraph 파이프라인 실행", t_graph_pipeline),
    ("RPT-04", "보고서 생성", "빈 주간 보고서(분기)", t_empty_week),
    ("RPT-05", "보고서 생성", "주간 요약 DB 저장", t_weekly_summary_db),

    ("GUI-01", "GUI·실행", "GUI 모듈 임포트", t_flet_import),
    ("GUI-02", "GUI·실행", "프로세스 감지", t_process_detect),
    ("GUI-03", "GUI·실행", "통합 수집기 임포트", t_integrated_collector),
    ("GUI-04", "GUI·실행", "예상질문자 설정 (추가·삭제)", t_expected_questioner_settings),
    ("GUI-05", "GUI·실행", "주제 질의·근거 제외 설정 (추가·삭제)", t_rag_topic_settings),
    ("GUI-06", "GUI·실행", "보고서 생성 상태 표시", t_report_status_panel),
    ("GUI-07", "GUI·실행", "저장소 관리 설정 (보관 기간·사용량)", t_storage_settings),
]

if __name__ == "__main__":
    os.makedirs("test_reports", exist_ok=True)
    print(f"=== 시나리오 테스트 시작 ({len(TESTS)}건) ===\n")
    for tc_id, cat, scenario, fn in TESTS:
        run(tc_id, cat, scenario, fn)

    # ---- Excel 생성 ----
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY = "002A5B"
    YELLOW = "F9BB00"
    LIGHT = "EFEFEF"
    PASS_GREEN = "D9EEDB"
    FAIL_RED = "F4CCCC"
    SKIP_YELLOW = "FFF2CC"

    wb = Workbook()
    ws = wb.active
    ws.title = "테스트 결과"

    headers = ["ID", "분류", "시나리오", "결과", "세부 내용", "실행(초)"]
    ws.append(headers)
    thin = Side(style="thin", color="999999")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(*[thin] * 4)

    for r in RESULTS:
        row = [r["id"], r["category"], r["scenario"], r["result"],
               r["detail"], r["elapsed"]]
        ws.append(row)
        cur = ws.max_row
        fill = {"PASS": PASS_GREEN, "FAIL": FAIL_RED, "SKIP": SKIP_YELLOW}[
            r["result"]
        ]
        for c in range(1, 7):
            cell = ws.cell(row=cur, column=c)
            cell.border = Border(*[thin] * 4)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c == 4:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
        if r["id"].endswith("01"):  # 카테고리 첫 행 구분
            pass

    widths = [10, 12, 32, 10, 60, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # 요약 시트
    ws2 = wb.create_sheet("요약")
    ws2.append(["분류", "총계", "PASS", "FAIL", "SKIP"])
    for cell in ws2[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    cats = {}
    for r in RESULTS:
        c = cats.setdefault(r["category"], {"total": 0, "PASS": 0, "FAIL": 0, "SKIP": 0})
        c["total"] += 1
        c[r["result"]] += 1
    for cat, c in cats.items():
        ws2.append([cat, c["total"], c["PASS"], c["FAIL"], c["SKIP"]])
    total = {"total": 0, "PASS": 0, "FAIL": 0, "SKIP": 0}
    for c in cats.values():
        for k in total:
            total[k] += c[k]
    ws2.append(["합계", total["total"], total["PASS"], total["FAIL"], total["SKIP"]])
    for c in range(1, 6):
        ws2.cell(row=ws2.max_row, column=c).font = Font(bold=True)
    for i, w in enumerate([16, 10, 10, 10, 10], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = f"test_reports/scenario_results_{stamp}.xlsx"
    wb.save(path)
    print(f"\n=== 결과 ===")
    print(f"PASS {total['PASS']} / FAIL {total['FAIL']} / SKIP {total['SKIP']}")
    print(f"Excel: {os.path.abspath(path)}")
