# -*- coding: utf-8 -*-
"""전체 기능 테스트 시나리오 실행 → Excel 결과 생성 (test_reports/)

카테고리:
  환경/인증 / 수집 / 데이터 가공 / AI 분석 / 보고서 생성 / GUI·실행
각 항목은 PASS / FAIL / SKIP 중 하나의 결과를 낸다.
"""

import json
import os
import re
import sys
import time
import traceback
from datetime import datetime, timedelta

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

RESULTS = []  # dict: id, category, scenario, result, detail, elapsed


def run(tc_id, category, scenario, fn):
    t0 = time.time()
    try:
        status, detail = fn()
    except Exception as e:
        status, detail = "FAIL", f"{type(e).__name__}: {e}"
        traceback.print_exc()
    RESULTS.append(
        {
            "id": tc_id,
            "category": category,
            "scenario": scenario,
            "result": status,
            "detail": detail,
            "elapsed": round(time.time() - t0, 1),
        }
    )
    print(f"[{status:4}] {tc_id} {scenario} — {detail[:80]}")


def ok(detail=""):
    return "PASS", detail


def fail(detail=""):
    return "FAIL", detail


def skip(detail=""):
    return "SKIP", detail


# ───────────────────────── 환경/인증 ─────────────────────────

def t_db_schema():
    from database import ActivityDatabase
    with ActivityDatabase() as db:
        rows = db.conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        ).fetchall()
        names = {r[0] for r in rows}
        need = {"activities", "weekly_summaries"}
        missing = need - names
        if missing:
            return fail(f"테이블 누락: {missing}")
        cnt = db.conn.execute("SELECT COUNT(*) FROM activities").fetchone()[0]
    return ok(f"테이블 정상, activities {cnt}건")


def t_litellm_config():
    from llm_summarizer import LLMSummarizer
    s = LLMSummarizer()
    if not s.enabled:
        return fail("config 비활성 또는 api_key 없음")
    return ok("config 로드 + enabled")


def t_embedding_call():
    from activity_clusterer import ActivityClusterer
    c = ActivityClusterer()
    if not c.enabled:
        return skip("LiteLLM config 없음")
    vecs = c.embed(["테스트 문장"])
    if not vecs:
        return fail("임베딩 호출 실패")
    return ok(f"임베딩 성공 dim={len(vecs[0])}")


def t_lcel_chat():
    from llm_summarizer import LLMSummarizer
    s = LLMSummarizer()
    if not s.enabled:
        return skip("LiteLLM config 없음")
    text = ("주간보고서 자동화 프로그램을 개발 중이다. " * 20)
    out = s.summarize(text, max_len=100, template="default")
    if not out or len(out) >= len(text):
        return fail(f"요약 미동작: {str(out)[:60]}")
    return ok(f"LCEL 체인 요약 성공 ({len(out)}자)")


def t_teams_token():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 파일 없음 (최초 로그인 필요)")
    with open("config/teams_graph_token.json", encoding="utf-8") as f:
        tok = json.load(f)
    has_rt = bool(tok.get("refresh_token"))
    if not has_rt:
        return fail("refresh_token 없음")
    return ok("토큰 파일 존재, refresh_token 보유")


def t_teams_auth_live():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 파일 없음")
    from teams_collector import get_access_token
    token = get_access_token()
    if not token:
        return fail("토큰 갱신/획득 실패")
    return ok("Graph 액세스 토큰 획득 (refresh 동작)")


# ───────────────────────── 수집 ─────────────────────────

def t_watcher_noise():
    from file_watcher import FileActivityHandler
    w = FileActivityHandler("logs/_test.log")
    cases = [
        ("C:/x/미확인 819675.crdownload", True),
        ("C:/x/(DDr^7D3.jpg", True),
        ("C:/x/D58F4000", True),
        ("C:/x/~$report.docx", True),
        ("C:/x/보고서.docx", False),
    ]
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    bad = [
        p for p, want in cases
        if bool(
            w.should_exclude(p)
            or rg._is_noise_activity(
                {"file_path": p, "source": "filesystem", "action": "modified",
                 "timestamp": "t"}
            )
        ) != want
    ]
    if bad:
        return fail(f"필터 오판: {bad}")
    return ok("노이즈 4종 + 정상파일 구분 정확")


def t_office_content_capture():
    """실제 오피스 파일 내용 추출 또는 DRM 폴백"""
    from file_watcher import FileActivityHandler
    w = FileActivityHandler("logs/_test.log")
    # xlsx DRM 폴백 경로 테스트: 실제 DRM 파일이 없으면 SKIP
    import glob
    drm_files = [
        p for p in glob.glob(r"C:\Users\SSG\Documents\**\*.xlsx", recursive=True)
        if os.path.exists(p)
    ][:3]
    if not drm_files:
        return skip("테스트용 xlsx 없음")
    res = w._capture_excel_content(drm_files[0])
    if res is None:
        return ok("비zip/잠김 → 조용히 None (정상 경로)")
    return ok(f"내용 추출/폴백 메시지: {str(res)[:60]}")


def t_browser_server():
    """브라우저 수신 서버 포트 상태 + POST 수신 경로"""
    import socket
    sock = socket.socket()
    sock.settimeout(2)
    try:
        sock.connect(("127.0.0.1", 5757))
        return ok("5757 포트 리스닝 중 (트래킹 실행 중)")
    except OSError:
        return skip("서버 미기동 — Start Tracking으로 실행되면 동작")
    finally:
        sock.close()


def t_ide_paths():
    from ide_collector import IDECollector
    c = IDECollector()
    paths = c._find_vscode_paths()
    if not paths:
        return skip("VS Code 경로 없음")
    return ok(f"VS Code 경로 {len(paths)}개 발견")


def t_teams_chat_collect():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from teams_collector import TeamsCollector
    col = TeamsCollector()
    acts = col.collect_all_teams_activity(days=1)
    return ok(f"Teams 활동 {len(acts)}건 수집")


def t_teams_calendar():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from teams_collector import TeamsCollector
    col = TeamsCollector()
    acts = col.collect_calendar_events(days=7)
    return ok(f"일정 {len(acts)}건 (취소 제외)")


def t_outlook_collect():
    from outlook_collector import OutlookCollector, _WIN32COM_AVAILABLE
    c = OutlookCollector()
    if _WIN32COM_AVAILABLE:
        connected = c.connect_outlook()
        if not connected:
            return skip("Outlook 미실행/COM 연결 실패")
        acts = c.collect_sent_email_activity(days=1)
        return ok(f"COM 경로 — 보낸 메일 {len(acts)}건")
    acts = c.collect_sent_email_activity(days=1)
    return ok(f"Graph 경로 — {len(acts)}건")


def t_confluence():
    cfg = "config/confluence_config.json"
    if not os.path.exists(cfg):
        return skip("Confluence config 없음")
    return ok("config 존재 (토큰 유효성은 수집 시 확인)")


def t_slack():
    cfg = "config/slack_config.json"
    if not os.path.exists(cfg):
        return skip("Slack 토큰 미발급 — 수집 비활성")
    return ok("config 존재")


def t_onenote():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from onenote_collector import OneNoteCollector
    col = OneNoteCollector()
    # 페이지 수준은 권한 부재 — 노트북/섹션 목록 수집 가능 여부 확인
    try:
        acts = col.collect_activity(days=400)
        return ok(f"노트북/섹션 변경 {len(acts)}건 수집 (페이지 단위는 권한 제한)")
    except Exception as e:
        return fail(str(e)[:80])


def t_sharepoint():
    if not os.path.exists("config/teams_graph_token.json"):
        return skip("토큰 없음")
    from sharepoint_collector import SharePointCollector
    col = SharePointCollector()
    try:
        # 버전 비교(다운로드+COM)는 느려서 여기선 수정 목록 수집만 확인 — 비교는 COL-13
        acts = col.collect_activity(days=7, with_changes=False)
        return ok(f"SharePoint/OneDrive 수정 파일 {len(acts)}건")
    except Exception as e:
        return fail(str(e)[:80])


def t_office_com_diff():
    """Office COM으로 xlsx 두 버전을 읽어 변경 줄을 뽑는지 (SharePoint 버전 비교 경로)"""
    if sys.platform != "win32":
        return skip("Windows 전용")
    import tempfile
    import openpyxl
    import office_reader
    tmp = tempfile.mkdtemp(prefix="wr_test_")
    paths = []
    for i, rows in enumerate([[("항목", "담당"), ("DB접속 암호화", "박성수")],
                              [("항목", "담당"), ("DB접속 암호화 (평문제거)", "박성수")]]):
        wb = openpyxl.Workbook()
        wb.active.title = "배포"
        for r in rows:
            wb.active.append(r)
        path = os.path.join(tmp, f"v{i}.xlsx")
        wb.save(path)
        paths.append(path)
    with office_reader.office_app("Excel") as app:
        if app is None:
            return skip("Excel 미설치")
        old, new = (office_reader.document_lines(app, "Excel", p) for p in paths)
    if old is None or new is None:
        return fail("COM으로 xlsx 읽기 실패")
    added, removed = office_reader.diff_lines(old, new)
    if added != ["[배포] DB접속 암호화 (평문제거) | 박성수"] or removed != ["[배포] DB접속 암호화 | 박성수"]:
        return fail(f"diff 불일치: +{added} -{removed}")
    return ok("COM 읽기 + 버전 diff 정상 (+1/-1줄)")


def t_file_watcher_watchdirs():
    cfg = "config/watch_config.json"
    if not os.path.exists(cfg):
        return fail("watch_config.json 없음")
    data = json.load(open(cfg, encoding="utf-8"))
    dirs = data.get("watch_paths", [])
    if not dirs:
        return fail("감시 폴더 0개")
    existing = [d for d in dirs if os.path.exists(d)]
    return ok(f"감시 폴더 {len(dirs)}개 (실존 {len(existing)})")


# ───────────────────────── 데이터 가공 ─────────────────────────

def t_meeting_dedupe():
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    acts = [
        {"timestamp": "2026-09-29T10:00:00", "action": "meeting",
         "file_path": "Meeting: 주간회의", "source": "outlook",
         "details": "{}"},
        {"timestamp": "2026-09-29T10:00:00", "action": "meeting",
         "file_path": "Teams Meeting: 주간회의", "source": "teams",
         "details": '{"subject": "주간회의"}'},
    ]
    out = rg._dedupe_meetings(acts)
    if len(out) == 1 and out[0]["source"] == "teams":
        return ok("동일 회의 중복 제거 + teams 우선")
    return fail(f"{len(out)}건 남음: {[a['source'] for a in out]}")


def t_transient_files():
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    acts = [
        {"timestamp": "t1", "action": "created", "file_path": "a.tmp", "source": "filesystem"},
        {"timestamp": "t2", "action": "modified", "file_path": "a.tmp", "source": "filesystem"},
        {"timestamp": "t3", "action": "deleted", "file_path": "a.tmp", "source": "filesystem"},
        {"timestamp": "t4", "action": "created", "file_path": "keep.docx", "source": "filesystem"},
    ]
    out = rg._drop_transient_files(acts)
    paths = {a["file_path"] for a in out}
    if paths == {"keep.docx"}:
        return ok("created+deleted 쌍 제거 정상")
    return fail(f"남은 경로: {paths}")


def t_path_normalize():
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    a = {"file_path": "C:/Users/SSG/test.txt"}
    out = rg._normalize_activity_path(a)
    if "/" in out["file_path"]:
        return fail(out["file_path"])
    return ok(f"정규화: {out['file_path']}")


def t_sensitive_url_filter():
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    bad = {
        "file_path": "https://login.example.com/callback?code=abc&state=x",
        "source": "browser", "action": "visited", "timestamp": "t",
    }
    if rg._is_noise_activity(bad):
        return ok("OAuth 콜백 URL 필터링됨")
    return fail("민감 URL 통과")


def t_summary_cache():
    from database import ActivityDatabase
    with ActivityDatabase() as db:
        try:
            cache = db.get_summary_cache("outlook")
            return ok(f"요약 캐시 {len(cache)}건 로드")
        except AttributeError:
            return fail("get_summary_cache 없음")


def t_shared_doc_split():
    """SharePoint: 내가 수정한 문서는 Excel, 파트원만 수정한 문서는 '공유 문서 변경'으로
    분리되고, 공유 문서 변경은 STEP 2 요약 입력에서 빠지는지"""
    from program_sections import build_program_sections, sections_digest, SHARED_DOCS

    def sp(name, editors):
        return {"timestamp": "2026-09-29T23:59:59", "source": "sharepoint", "action": "modified",
                "file_path": f"SharePoint 파일: {name}", "file_type": "sharepoint",
                "details": json.dumps({"name": name, "editors": editors, "summary": f"{name} 변경"})}

    me = "홍길동(POS서버) - 팀A"
    secs = build_program_sections(
        [sp("내파일.xlsx", ["홍길동(POS서버) - 팀B", "김철수(POS) - 팀A"]),
         sp("팀파일.xlsx", ["김철수(POS) - 팀A"])],
        me_name=me,
    )
    by_name = {s["name"]: [i["title"] for i in s["items"]] for s in secs}
    if by_name.get("Excel") != ["내파일.xlsx (SharePoint)"] or by_name.get(SHARED_DOCS) != ["팀파일.xlsx (SharePoint)"]:
        return fail(f"분류 오류: {by_name}")
    if "팀파일" in sections_digest(secs):
        return fail("공유 문서 변경이 STEP 2 입력에 포함됨")
    return ok("내 수정→Excel, 파트원 수정→공유 문서 변경, STEP 2 제외")


def t_vector_store():
    """VectorDB(Qdrant 로컬): 저장·재조회·유사도 검색·날짜/소스/발신자 필터 (임시 폴더, 4차원 벡터)"""
    import json as _json, shutil, tempfile
    from vector_store import ActivityVectorStore
    tmp = tempfile.mkdtemp(prefix="wr_qdrant_")
    try:
        store = ActivityVectorStore(path=tmp, vector_size=4)

        def act(ts, source, path, sender=""):
            return {"timestamp": ts, "action": "message", "file_path": path, "source": source,
                    "details": _json.dumps({"sender": sender}, ensure_ascii=False)}

        a1 = act("2026-09-28T10:00:00", "teams", "배포 일정", "김영호(POS) - 팀")
        a2 = act("2026-09-29T10:00:00", "teams", "장애 대응", "홍길동(POS) - 팀")
        a3 = act("2026-09-20T10:00:00", "outlook", "배포 공지")
        store.upsert([(a1, "배포 일정", [1, 0, 0, 0]), (a2, "장애 대응", [0, 1, 0, 0]), (a3, "배포 공지", [0.9, 0.1, 0, 0])])
        store.upsert([(a1, "배포 일정", [1, 0, 0, 0])])  # 같은 활동 재저장 → 덮어쓰기
        if store.count() != 3:
            return fail(f"건수 {store.count()} (재저장 중복)")
        from database import ActivityDatabase
        got = store.get_vectors([ActivityDatabase.activity_key(a1)])
        if list(got.values())[0][0] != "배포 일정":
            return fail("재조회 실패")
        top = store.search([1, 0, 0, 0], limit=2)
        if [h["text"] for h in top] != ["배포 일정", "배포 공지"]:
            return fail(f"유사도 순서 오류: {[h['text'] for h in top]}")
        if [h["text"] for h in store.search([1, 0, 0, 0], date_from="2026-09-25", date_to="2026-09-28")] != ["배포 일정"]:
            return fail("날짜 필터 오류")
        if {h["source"] for h in store.search([1, 0, 0, 0], sources=["outlook"])} != {"outlook"}:
            return fail("소스 필터 오류")
        if [h["text"] for h in store.search([1, 0, 0, 0], sender_contains="홍길동")] != ["장애 대응"]:
            return fail("발신자 필터 오류")
        return ok("저장·덮어쓰기·재조회·유사도 순서·날짜/소스/발신자 필터 정상")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def t_split_sequence():
    """요약 줄 안의 번호 목록 분리 — IP·버전·날짜는 번호로 오인하지 않는지"""
    from program_sections import split_sequence
    got = split_sequence("개선 요청: 1. 데이터 수집 문제 2. 중복 제거 3. 요약 확대")
    if got != {"head": "개선 요청:", "items": ["1. 데이터 수집 문제", "2. 중복 제거", "3. 요약 확대"]}:
        return fail(f"분리 오류: {got}")
    for text in ("개발 DB 10.253.42.50 1526 접속", "v3.0 배포 후 2. 항목만", "09/29(화) 1. 단독 항목"):
        if split_sequence(text)["items"]:
            return fail(f"오탐: {text}")
    return ok("번호 목록 분리 + IP/버전/단독 번호 오탐 없음")


# ───────────────────────── AI 분석 ─────────────────────────

def t_llm_templates():
    from llm_summarizer import LLMSummarizer, PROMPT_TEMPLATES
    need = {"default", "file", "teams", "mail", "topic"}
    missing = need - set(PROMPT_TEMPLATES)
    if missing:
        return fail(f"템플릿 누락: {missing}")
    return ok("소스별 프롬프트 템플릿 5종 존재")


def t_cluster_topics():
    from report_generator import ReportGenerator
    from database import ActivityDatabase, get_week_start_end
    start, end = get_week_start_end()
    with ActivityDatabase() as db:
        acts = db.get_activities_by_date_range(start, end)
    rg = ReportGenerator()
    topics = rg._cluster_topics(acts)
    if not topics:
        return fail("주제 클러스터링 결과 없음")
    names = [t["name"] for t in topics[:3]]
    return ok(f"{len(topics)}개 주제 — 예: {', '.join(names)}")


def t_langchain_present():
    try:
        import langchain_core, langchain_openai, langgraph  # noqa
        return ok("langchain-core/openai/langgraph 임포트 정상")
    except ImportError as e:
        return fail(str(e))


# ───────────────────────── 보고서 생성 ─────────────────────────

def t_markdown_report():
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    data = rg.collect_weekly_data()
    md = rg.generate_markdown(data)
    required = ["STEP 1. 프로그램별 상세 활동 내역", "STEP 2. 업무 요약", "이번 주 핵심 요약"]
    missing = [r for r in required if r not in md]
    if missing:
        return fail(f"섹션 누락: {missing}")
    return ok(f"마크다운 생성 ({len(md)}자)")


def t_word_report():
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    data = rg.collect_weekly_data()
    path = "test_reports/_test_report.docx"
    os.makedirs("test_reports", exist_ok=True)
    try:
        rg.generate_word(data, path)
        size = os.path.getsize(path)
        if size < 1000:
            return fail(f"파일 너무 작음 {size}B")
        return ok(f"docx 생성 {size//1024}KB")
    finally:
        if os.path.exists(path):
            os.remove(path)


def t_graph_pipeline():
    # test_reports에 저장 — reports/의 이번 주 보고서(사용자가 편집한 파일)를 덮어쓰지 않도록
    from report_pipeline import WeeklyReportPipeline
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    result = WeeklyReportPipeline(rg, output_dir="test_reports").run(formats=("markdown",))
    paths = result["file_paths"]
    if "markdown" not in paths or not os.path.exists(paths["markdown"]):
        return fail("마크다운 경로 없음")
    return ok(f"LangGraph 경로 생성: {os.path.basename(paths['markdown'])}")


def t_empty_week():
    """활동 없는 주에도 보고서가 깨지지 않는지 (analyze 스킵 분기)"""
    from report_pipeline import WeeklyReportPipeline
    from report_generator import ReportGenerator
    rg = ReportGenerator()
    pipe = WeeklyReportPipeline(rg, output_dir="test_reports")
    result = pipe.run(week_start="2026-01-01", week_end="2026-01-07",
                      formats=("markdown",))
    wd = result["weekly_data"]
    if wd["has_activities"]:
        return fail("빈 주인데 활동 있음")
    return ok("빈 주 보고서 생성 + analyze 스킵 정상")


def t_weekly_summary_db():
    from database import ActivityDatabase
    with ActivityDatabase() as db:
        rows = db.conn.execute(
            "SELECT COUNT(*) FROM weekly_summaries"
        ).fetchone()[0]
    if rows == 0:
        return fail("weekly_summaries 비어있음")
    return ok(f"주간 요약 {rows}건 저장됨")


# ───────────────────────── GUI·실행 ─────────────────────────

def t_flet_import():
    try:
        import modern_gui  # noqa
        return ok("modern_gui 임포트 정상 (flet+win32 가드)")
    except ImportError as e:
        return fail(str(e))


def t_process_detect():
    from modern_gui import WeeklyPulseApp
    for proc in ["chrome.exe", "winword.exe", "Code.exe", "powershell.exe"]:
        if WeeklyPulseApp.is_app_running([proc]):
            return ok(f"{proc} 감지됨")
    return ok("감지 대상 프로세스 없음 (로직 실행 정상)")


def t_integrated_collector():
    try:
        from integrated_collector import IntegratedCollector
        return ok("IntegratedCollector 임포트 정상")
    except ImportError as e:
        return fail(str(e))


def t_expected_questioner_settings():
    """Settings > Teams 예상질문자 추가·중복·빈값·삭제 (실제 설정 파일은 저장하지 않음)"""
    from types import SimpleNamespace
    from teams_collector import DEFAULT_SETTINGS
    from modern_gui import WeeklyPulseApp
    if DEFAULT_SETTINGS.get("expected_questioners") != []:
        return fail("기본값 expected_questioners 누락")
    app = WeeklyPulseApp.__new__(WeeklyPulseApp)
    app.page = SimpleNamespace(window=SimpleNamespace(width=1400), update=lambda: None)
    app.show_snack = lambda msg: None
    app._save_teams_settings = lambda: None
    app._create_teams_settings_section()
    app.teams_settings["expected_questioners"] = []
    for text in ("  김영호 ", "팀장", "김영호", "   "):
        app.questioner_input.value = text
        app._add_expected_questioner(None)
    app._remove_expected_questioner("팀장")
    names = app.teams_settings["expected_questioners"]
    if names != ["김영호"] or len(app.questioner_chips.controls) != 1:
        return fail(f"결과 불일치: {names}")
    return ok("추가·공백정리·중복/빈값 거부·삭제 정상")


def t_rag_topic_settings():
    """Settings > 보고서 설정 주제 질의 추가·중복·삭제·검색 범위 (실제 설정 파일은 저장하지 않음)"""
    from types import SimpleNamespace
    from modern_gui import WeeklyPulseApp
    app = WeeklyPulseApp.__new__(WeeklyPulseApp)
    app.page = SimpleNamespace(window=SimpleNamespace(width=1400), update=lambda: None)
    app.show_snack = lambda msg: None
    app._save_report_settings = lambda: None
    app._create_report_settings_section()
    app.report_settings["rag_topics"] = ["장애 대응"]
    for text in (" 손익·예산 ", "장애 대응", "  "):
        app.rag_topic_input.value = text
        app._add_rag_topic(None)
    app._remove_rag_topic("장애 대응")
    app._rag_past_weeks_changed("abc")
    app._rag_past_weeks_changed("6")
    topics = app.report_settings["rag_topics"]
    if topics != ["손익·예산"] or len(app.rag_topic_chips.controls) != 1:
        return fail(f"결과 불일치: {topics}")
    if app.report_settings["rag_past_weeks"] != 6:
        return fail(f"검색 범위 불일치: {app.report_settings['rag_past_weeks']}")
    return ok("주제 추가·중복/빈값 거부·삭제, 검색 범위 숫자 검증 정상")


def t_rag_citation_check():
    """RAG 출력 검증: 인용 없는 문장·없는 근거 번호 문장 제거, 인용 과다 축약, NONE 처리"""
    from rag import parse_cited_bullets
    output = (
        "- 거래저장 실패 장애를 복구함 [1][2]\n"
        "- 근거 없는 추측 문장\n"
        "- 존재하지 않는 근거를 단 문장 [9]\n"
        "서론 문장 [1]\n"
        "- 근거를 많이 단 문장 [1][2][3][4]\n"
        "- (지난 기록) 지난달 점검과 이어짐 [3]"
    )
    bullets = parse_cited_bullets(output, evidence_count=4)
    texts = [t for t, _ in bullets]
    if len(bullets) != 3 or bullets[1][1] != [1, 2, 3] or "[4]" in texts[1]:
        return fail(f"검증 결과 불일치: {bullets}")
    if parse_cited_bullets("NONE", 3) != [] or parse_cited_bullets("", 3) != []:
        return fail("NONE/빈 출력 처리 실패")
    return ok("근거 없는 문장 2건 제거, 인용 4개→3개 축약, NONE 처리 정상")


def t_rag_sections():
    """실데이터 RAG: 이번 주 활동 → 주제별 근거 인용 요약 → Markdown 'STEP 2 주제 질의 요약'"""
    from report_generator import ReportGenerator
    from rag import TopicQueryRAG
    rg = ReportGenerator()
    week_start, week_end = rg._resolve_week_range(None, None)
    acts = rg.fetch_week_activities(week_start, week_end)
    rag = TopicQueryRAG(rg)
    if not rag.enabled:
        return skip("LiteLLM 설정 없음")
    if not acts:
        return skip("이번 주 활동 없음")
    sections = rag.build_sections(acts, week_start, week_end)
    if not sections:
        return skip("설정한 주제와 관련된 이번 주 기록 없음")
    for s in sections:
        numbers = {e["n"] for e in s["evidence"]}
        cited = {int(n) for b in s["bullets"] for n in re.findall(r"\[(\d+)\]", b)}
        if not s["bullets"] or not cited or not cited <= numbers:
            return fail(f"[{s['topic']}] 인용·근거 번호 불일치: {cited} / {numbers}")
    data = rg.compose_weekly_data(week_start, week_end, acts, {"rag_sections": sections})
    md = rg.generate_markdown(data)
    if "### 주제 질의 요약" not in md or md.index("### 주제 질의 요약") < md.index("## STEP 2"):
        return fail("Markdown STEP 2에 주제 질의 요약 없음")
    return ok(f"{len(sections)}개 주제 — " + ", ".join(s["topic"] for s in sections))


# ───────────────────────── 실행 ─────────────────────────

TESTS = [
    ("ENV-01", "환경/인증", "DB 스키마 및 연결", t_db_schema),
    ("ENV-02", "환경/인증", "LiteLLM 프록시 설정 로드", t_litellm_config),
    ("ENV-03", "환경/인증", "임베딩 API 호출", t_embedding_call),
    ("ENV-04", "환경/인증", "LLM 채팅(LCEL 체인) 호출", t_lcel_chat),
    ("ENV-05", "환경/인증", "Teams 토큰 파일 존재", t_teams_token),
    ("ENV-06", "환경/인증", "Teams 토큰 갱신(라이브)", t_teams_auth_live),

    ("COL-01", "수집", "파일 감시 노이즈 필터", t_watcher_noise),
    ("COL-02", "수집", "오피스 내용 캡처/DRM 폴백", t_office_content_capture),
    ("COL-03", "수집", "브라우저 수신 서버(5757)", t_browser_server),
    ("COL-04", "수집", "IDE 경로 탐색", t_ide_paths),
    ("COL-05", "수집", "Teams 채팅 수집(라이브)", t_teams_chat_collect),
    ("COL-06", "수집", "Teams 일정 수집(라이브)", t_teams_calendar),
    ("COL-07", "수집", "Outlook 수집(COM/Graph)", t_outlook_collect),
    ("COL-08", "수집", "Confluence 설정 존재", t_confluence),
    ("COL-09", "수집", "Slack 설정 존재", t_slack),
    ("COL-10", "수집", "감시 폴더 설정", t_file_watcher_watchdirs),
    ("COL-11", "수집", "OneNote 수집(라이브)", t_onenote),
    ("COL-12", "수집", "SharePoint/OneDrive 수집(라이브)", t_sharepoint),
    ("COL-13", "수집", "Office COM 읽기 + 버전 diff", t_office_com_diff),

    ("PROC-01", "데이터 가공", "회의 중복 제거(teams 우선)", t_meeting_dedupe),
    ("PROC-02", "데이터 가공", "일시 파일(created+deleted) 제거", t_transient_files),
    ("PROC-03", "데이터 가공", "경로 정규화", t_path_normalize),
    ("PROC-04", "데이터 가공", "민감 URL 필터(OAuth 콜백)", t_sensitive_url_filter),
    ("PROC-05", "데이터 가공", "요약 캐시 조회", t_summary_cache),
    ("PROC-06", "데이터 가공", "공유 문서 변경 분리(내 수정/파트원 수정)", t_shared_doc_split),
    ("PROC-07", "데이터 가공", "요약 번호 목록 줄바꿈 분리", t_split_sequence),
    ("PROC-08", "데이터 가공", "VectorDB 저장·필터 검색 (Qdrant)", t_vector_store),

    ("AI-01", "AI 분석", "소스별 프롬프트 템플릿", t_llm_templates),
    ("AI-02", "AI 분석", "임베딩 클러스터링 → 주제", t_cluster_topics),
    ("AI-03", "AI 분석", "LangChain/LangGraph 임포트", t_langchain_present),
    ("AI-04", "AI 분석", "RAG 근거 인용 검증 (하네스)", t_rag_citation_check),
    ("AI-05", "AI 분석", "RAG 주제 질의 섹션 (실데이터)", t_rag_sections),

    ("RPT-01", "보고서 생성", "Markdown 섹션 완전성", t_markdown_report),
    ("RPT-02", "보고서 생성", "Word(docx) 생성", t_word_report),
    ("RPT-03", "보고서 생성", "LangGraph 파이프라인 실행", t_graph_pipeline),
    ("RPT-04", "보고서 생성", "빈 주간 보고서(분기)", t_empty_week),
    ("RPT-05", "보고서 생성", "주간 요약 DB 저장", t_weekly_summary_db),

    ("GUI-01", "GUI·실행", "GUI 모듈 임포트", t_flet_import),
    ("GUI-02", "GUI·실행", "프로세스 감지", t_process_detect),
    ("GUI-03", "GUI·실행", "통합 수집기 임포트", t_integrated_collector),
    ("GUI-04", "GUI·실행", "예상질문자 설정 (추가·삭제)", t_expected_questioner_settings),
    ("GUI-05", "GUI·실행", "주제 질의 설정 (추가·삭제)", t_rag_topic_settings),
]

if __name__ == "__main__":
    os.makedirs("test_reports", exist_ok=True)
    print(f"=== 시나리오 테스트 시작 ({len(TESTS)}건) ===\n")
    for tc_id, cat, scenario, fn in TESTS:
        run(tc_id, cat, scenario, fn)

    # ---- Excel 생성 ----
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY = "002A5B"
    YELLOW = "F9BB00"
    LIGHT = "EFEFEF"
    PASS_GREEN = "D9EEDB"
    FAIL_RED = "F4CCCC"
    SKIP_YELLOW = "FFF2CC"

    wb = Workbook()
    ws = wb.active
    ws.title = "테스트 결과"

    headers = ["ID", "분류", "시나리오", "결과", "세부 내용", "실행(초)"]
    ws.append(headers)
    thin = Side(style="thin", color="999999")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(*[thin] * 4)

    for r in RESULTS:
        row = [r["id"], r["category"], r["scenario"], r["result"],
               r["detail"], r["elapsed"]]
        ws.append(row)
        cur = ws.max_row
        fill = {"PASS": PASS_GREEN, "FAIL": FAIL_RED, "SKIP": SKIP_YELLOW}[
            r["result"]
        ]
        for c in range(1, 7):
            cell = ws.cell(row=cur, column=c)
            cell.border = Border(*[thin] * 4)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c == 4:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
        if r["id"].endswith("01"):  # 카테고리 첫 행 구분
            pass

    widths = [10, 12, 32, 10, 60, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # 요약 시트
    ws2 = wb.create_sheet("요약")
    ws2.append(["분류", "총계", "PASS", "FAIL", "SKIP"])
    for cell in ws2[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    cats = {}
    for r in RESULTS:
        c = cats.setdefault(r["category"], {"total": 0, "PASS": 0, "FAIL": 0, "SKIP": 0})
        c["total"] += 1
        c[r["result"]] += 1
    for cat, c in cats.items():
        ws2.append([cat, c["total"], c["PASS"], c["FAIL"], c["SKIP"]])
    total = {"total": 0, "PASS": 0, "FAIL": 0, "SKIP": 0}
    for c in cats.values():
        for k in total:
            total[k] += c[k]
    ws2.append(["합계", total["total"], total["PASS"], total["FAIL"], total["SKIP"]])
    for c in range(1, 6):
        ws2.cell(row=ws2.max_row, column=c).font = Font(bold=True)
    for i, w in enumerate([16, 10, 10, 10, 10], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = f"test_reports/scenario_results_{stamp}.xlsx"
    wb.save(path)
    print(f"\n=== 결과 ===")
    print(f"PASS {total['PASS']} / FAIL {total['FAIL']} / SKIP {total['SKIP']}")
    print(f"Excel: {os.path.abspath(path)}")
